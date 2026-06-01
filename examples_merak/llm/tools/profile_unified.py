#!/usr/bin/env python3
"""All-in-one Qwen profiler + Excel generator.

Single-file implementation (no importlib/reference to other local profile/gen scripts).
Supports both Qwen3.5 and Qwen3-Next configs, with summary/next styles.

Examples
--------
Profile summary:
  python examples/qwen3.5/profile_unified.py profile --style summary --config weights/Qwen3.5-27B/config.json --no-log

Profile next:
  python examples/qwen3.5/profile_unified.py profile --style next --config weights/Qwen3.5-27B/config.json --no-log

Profile util:
  python examples/qwen3.5/profile_unified.py profile --style util --config weights/Qwen3.5-27B/config.json --bandwidth 272 --compute-utilization 0.5 --bandwidth-utilization 0.7 --no-log

Excel summary:
  python examples/qwen3.5/profile_unified.py excel --style summary --config weights/Qwen3.5-27B/config.json

Excel next:
  python examples/qwen3.5/profile_unified.py excel --style next --config weights/Qwen3.5-27B/config.json

Excel util:
  python examples/qwen3.5/profile_unified.py excel --style util --config weights/Qwen3.5-27B/config.json --bandwidth 272 --compute-utilization 0.5 --bandwidth-utilization 0.7
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from contextlib import redirect_stdout
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence


BYTES_PER_DTYPE = {
    "fp32": 4.0,
    "float32": 4.0,
    "fp16": 2.0,
    "float16": 2.0,
    "bf16": 2.0,
    "bfloat16": 2.0,
    "int16": 2.0,
    "sefp16": 1.0,
    "int8": 1.0,
    "int4": 0.5,
    "ssfpw4": 4.5 / 8.0,
}

SIZE_HINT_RE = re.compile(r"\d+(?:\.\d+)?B(?:-A\d+(?:\.\d+)?B)?", re.IGNORECASE)
INVALID_SHEET_CHARS = set('[]:*?/\\')


def normalize_dtype(dtype: str) -> str:
    return dtype.strip().lower()


def resolve_dtype_bytes(dtype: str, bytes_override: Optional[float]) -> float:
    if bytes_override is not None:
        if bytes_override <= 0:
            raise ValueError("bytes override must be > 0")
        return float(bytes_override)
    key = normalize_dtype(dtype)
    if key not in BYTES_PER_DTYPE:
        raise ValueError(f"Unsupported dtype '{dtype}'. Supported keys: {sorted(BYTES_PER_DTYPE)}")
    return BYTES_PER_DTYPE[key]


def fmt_num(n: float, unit: str = "", si: bool = False) -> str:
    if si:
        prefixes = [(1e15, "P"), (1e12, "T"), (1e9, "G"), (1e6, "M"), (1e3, "K")]
    else:
        prefixes = [(1e15, "P"), (1e12, "T"), (1e9, "B"), (1e6, "M"), (1e3, "K")]
    for threshold, suffix in prefixes:
        if abs(n) >= threshold:
            return f"{n / threshold:.2f}{suffix}{unit}"
    return f"{n:.2f}{unit}"


def fmt_flops(n: float) -> str:
    return fmt_num(n, " FLOPs", si=True)


def fmt_bytes(n: float) -> str:
    for threshold, suffix in [(1e12, "TB"), (1e9, "GB"), (1e6, "MB"), (1e3, "KB")]:
        if abs(n) >= threshold:
            return f"{n / threshold:.2f} {suffix}"
    return f"{n:.0f} B"


def fmt_context(ctx: int) -> str:
    if ctx % (1024 * 1024) == 0 and ctx >= 1024 * 1024:
        return f"{ctx // (1024 * 1024)}M"
    if ctx % 1024 == 0:
        return f"{ctx // 1024}K"
    return str(ctx)


def fmt_context_k(ctx: int) -> str:
    value = ctx / 1024.0
    if value.is_integer():
        return str(int(value))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def fmt_time(ms: float) -> str:
    if ms >= 1000:
        return f"{ms / 1000:.3f} s"
    return f"{ms:.4f} ms"


def sanitize_model_name(name: str) -> str:
    cleaned = name.strip().replace("\\", "/").replace("/", "_").replace(" ", "_")
    while "__" in cleaned:
        cleaned = cleaned.replace("__", "_")
    return cleaned.strip("_") or "model"


def infer_model_name(config_path: str) -> str:
    path = Path(config_path)
    path_name = sanitize_model_name(path.parent.name or path.stem or "model")
    if SIZE_HINT_RE.search(path_name):
        return path_name

    try:
        with open(path, "r", encoding="utf-8") as f:
            raw = json.load(f)
        cfg = raw.get("text_config", raw) if isinstance(raw, dict) else {}
        candidates: List[Any] = []
        if isinstance(raw, dict):
            candidates.extend([raw.get("_name_or_path"), raw.get("name_or_path")])
        if isinstance(cfg, dict):
            candidates.extend([cfg.get("_name_or_path"), cfg.get("name_or_path")])
        for cand in candidates:
            if isinstance(cand, str) and cand.strip():
                name = sanitize_model_name(cand)
                if SIZE_HINT_RE.search(name):
                    return name
    except Exception:
        pass

    return path_name


def make_output_path(model_name: str, style: str, kind: str, output: Optional[str]) -> str:
    if output:
        out = Path(output)
        if output.endswith("/") or (out.exists() and out.is_dir()):
            out.mkdir(parents=True, exist_ok=True)
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            return str(out / f"{model_name}_{kind}_{style}_{ts}.xlsx")
        parent = out.parent
        if str(parent) and str(parent) != ".":
            parent.mkdir(parents=True, exist_ok=True)
        return output

    os.makedirs("output", exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    suffix = ".xlsx" if kind == "profile" else ".log"
    return f"output/{model_name}_{kind}_{style}_{ts}{suffix}"


def safe_sheet_name(model_name: str, suffix: str) -> str:
    base = "".join("_" if ch in INVALID_SHEET_CHARS else ch for ch in model_name).strip() or "Model"
    title = f"{base}-{suffix}"
    return title[:31]


def brief_precision_label(dtype: str, prefix: str) -> str:
    key = normalize_dtype(dtype)
    if "4" in key:
        return f"{prefix}4"
    if "8" in key:
        return f"{prefix}8"
    if "16" in key or key in {"bf16", "bfloat16", "fp16", "float16", "sefp16"}:
        return f"{prefix}16"
    if "32" in key or key in {"fp32", "float32"}:
        return f"{prefix}32"
    return f"{prefix}?"


@dataclass
class ModelConfig:
    hidden_size: int = 4096
    num_hidden_layers: int = 32
    head_dim: int = 256
    vocab_size: int = 248320

    num_attention_heads: int = 16
    num_key_value_heads: int = 4

    linear_key_head_dim: int = 128
    linear_num_key_heads: int = 16
    linear_value_head_dim: int = 128
    linear_num_value_heads: int = 32
    linear_conv_kernel_dim: int = 4

    intermediate_size: int = 12288
    num_experts: int = 0
    num_experts_per_tok: int = 0
    moe_intermediate_size: int = 0
    shared_expert_intermediate_size: int = 0
    full_attention_interval: int = 4
    layer_types: Optional[List[str]] = None

    @classmethod
    def from_json(cls, path: str) -> "ModelConfig":
        with open(path, "r", encoding="utf-8") as f:
            raw = json.load(f)
        cfg = raw.get("text_config", raw)
        valid = set(cls.__dataclass_fields__.keys())
        payload = {k: v for k, v in cfg.items() if k in valid}
        return cls(**payload)

    def __post_init__(self) -> None:
        if not self.layer_types or len(self.layer_types) != self.num_hidden_layers:
            self.layer_types = [
                "full_attention" if (i + 1) % self.full_attention_interval == 0 else "linear_attention"
                for i in range(self.num_hidden_layers)
            ]
        if self.moe_intermediate_size <= 0:
            self.moe_intermediate_size = self.intermediate_size

    @property
    def is_moe(self) -> bool:
        return self.num_experts > 0 and self.num_experts_per_tok > 0

    @property
    def num_full_attn_layers(self) -> int:
        return sum(1 for t in self.layer_types or [] if t == "full_attention")

    @property
    def num_linear_attn_layers(self) -> int:
        return self.num_hidden_layers - self.num_full_attn_layers

    @property
    def lin_key_dim(self) -> int:
        return self.linear_num_key_heads * self.linear_key_head_dim

    @property
    def lin_value_dim(self) -> int:
        return self.linear_num_value_heads * self.linear_value_head_dim

    @property
    def lin_conv_dim(self) -> int:
        return self.lin_key_dim * 2 + self.lin_value_dim

    @property
    def full_q_out_dim(self) -> int:
        return self.num_attention_heads * self.head_dim

    @property
    def full_qz_out_dim(self) -> int:
        return self.num_attention_heads * self.head_dim * 2

    @property
    def full_kv_out_dim(self) -> int:
        return self.num_key_value_heads * self.head_dim


class QwenProfiler:
    def __init__(
        self,
        config: ModelConfig,
        bytes_per_weight: float,
        bytes_per_activation: float,
        include_activation_io: bool = False,
    ):
        self.c = config
        self.h = config.hidden_size
        self.bpw = float(bytes_per_weight)
        self.bpa = float(bytes_per_activation)
        self.include_activation_io = bool(include_activation_io)

    @staticmethod
    def _linear_flops(in_dim: int, out_dim: int, tokens: int = 1) -> int:
        return 2 * in_dim * out_dim * tokens

    def params_embedding(self) -> Dict[str, int]:
        return {"embed_tokens": self.c.vocab_size * self.h}

    def params_lm_head(self) -> Dict[str, int]:
        return {"lm_head": self.h * self.c.vocab_size}

    def params_final_norm(self) -> Dict[str, int]:
        return {"final_norm": self.h}

    def params_layer_norms(self) -> Dict[str, int]:
        return {"input_layernorm": self.h, "post_attn_layernorm": self.h}

    def params_full_attn(self) -> Dict[str, int]:
        c, h = self.c, self.h
        return {
            "Q_proj": h * c.full_q_out_dim,
            "Z_proj": h * c.full_q_out_dim,
            "K_proj": h * c.full_kv_out_dim,
            "V_proj": h * c.full_kv_out_dim,
            "O_proj": c.full_q_out_dim * h,
            "q_norm": c.head_dim,
            "k_norm": c.head_dim,
        }

    def params_linear_attn(self) -> Dict[str, int]:
        c, h = self.c, self.h
        return {
            "Q_proj": h * c.lin_key_dim,
            "K_proj": h * c.lin_key_dim,
            "V_proj": h * c.lin_value_dim,
            "Z_proj": h * c.lin_value_dim,
            "A_proj": h * c.linear_num_value_heads,
            "G_proj": h * c.linear_num_value_heads,
            "O_proj": c.lin_value_dim * h,
            "conv1d": c.lin_conv_dim * c.linear_conv_kernel_dim,
            "dt_bias": c.linear_num_value_heads,
            "A_log": c.linear_num_value_heads,
            "gated_norm": c.linear_value_head_dim,
        }

    def params_dense_ffn(self) -> Dict[str, int]:
        c, h = self.c, self.h
        return {
            "gate_proj": h * c.intermediate_size,
            "up_proj": h * c.intermediate_size,
            "down_proj": c.intermediate_size * h,
        }

    def params_moe(self, activated_only: bool = False) -> Dict[str, int]:
        c, h = self.c, self.h
        n_exp = c.num_experts_per_tok if activated_only else c.num_experts
        expert_mid = c.moe_intermediate_size
        shared_mid = c.shared_expert_intermediate_size
        expert_params = 3 * h * expert_mid
        shared_params = 3 * h * shared_mid
        return {
            "router": h * c.num_experts,
            "shared_expert_gate": h if shared_mid > 0 else 0,
            f"experts(x{n_exp})": expert_params * n_exp,
            "shared_expert": shared_params,
        }

    def params_mlp(self, activated_only: bool = False) -> Dict[str, int]:
        if self.c.is_moe:
            return self.params_moe(activated_only=activated_only)
        return self.params_dense_ffn()

    def total_params(self) -> int:
        c = self.c
        mlp_total = sum(self.params_mlp(activated_only=False).values())
        per_full = sum(self.params_full_attn().values()) + sum(self.params_layer_norms().values()) + mlp_total
        per_linear = sum(self.params_linear_attn().values()) + sum(self.params_layer_norms().values()) + mlp_total
        return (
            sum(self.params_embedding().values())
            + sum(self.params_lm_head().values())
            + sum(self.params_final_norm().values())
            + per_full * c.num_full_attn_layers
            + per_linear * c.num_linear_attn_layers
        )

    def activated_params(self) -> int:
        c = self.c
        if not c.is_moe:
            return self.total_params()
        mlp_active = sum(self.params_mlp(activated_only=True).values())
        per_full = sum(self.params_full_attn().values()) + sum(self.params_layer_norms().values()) + mlp_active
        per_linear = sum(self.params_linear_attn().values()) + sum(self.params_layer_norms().values()) + mlp_active
        return (
            sum(self.params_embedding().values())
            + sum(self.params_lm_head().values())
            + sum(self.params_final_norm().values())
            + per_full * c.num_full_attn_layers
            + per_linear * c.num_linear_attn_layers
        )

    def flops_full_attn_proj(self, tokens: int = 1) -> Dict[str, int]:
        c, h = self.c, self.h
        return {
            "Q_proj": self._linear_flops(h, c.full_q_out_dim, tokens),
            "Z_proj": self._linear_flops(h, c.full_q_out_dim, tokens),
            "K_proj": self._linear_flops(h, c.full_kv_out_dim, tokens),
            "V_proj": self._linear_flops(h, c.full_kv_out_dim, tokens),
            "O_proj": self._linear_flops(c.full_q_out_dim, h, tokens),
        }

    def flops_full_attn_score(self, tokens: int, context_len: int) -> Dict[str, int]:
        c = self.c
        n_heads = c.num_attention_heads
        d = c.head_dim
        return {
            "QK^T": 2 * n_heads * d * tokens * context_len,
            "softmax": 5 * n_heads * tokens * context_len,
            "score@V": 2 * n_heads * d * tokens * context_len,
        }

    def flops_linear_attn_proj(self, tokens: int = 1) -> Dict[str, int]:
        c, h = self.c, self.h
        return {
            "Q_proj": self._linear_flops(h, c.lin_key_dim, tokens),
            "K_proj": self._linear_flops(h, c.lin_key_dim, tokens),
            "V_proj": self._linear_flops(h, c.lin_value_dim, tokens),
            "Z_proj": self._linear_flops(h, c.lin_value_dim, tokens),
            "A_proj": self._linear_flops(h, c.linear_num_value_heads, tokens),
            "G_proj": self._linear_flops(h, c.linear_num_value_heads, tokens),
            "O_proj": self._linear_flops(c.lin_value_dim, h, tokens),
        }

    def flops_linear_attn_conv(self, tokens: int = 1) -> Dict[str, int]:
        c = self.c
        return {
            "conv1d": 2 * c.lin_conv_dim * c.linear_conv_kernel_dim * tokens,
            "silu": 4 * c.lin_conv_dim * tokens,
        }

    def flops_linear_attn_recurrent(self, tokens: int = 1) -> Dict[str, int]:
        c = self.c
        state_size = c.linear_num_value_heads * c.linear_key_head_dim * c.linear_value_head_dim
        return {
            "state_gate": state_size * tokens,
            "state_update": 2 * state_size * tokens,
            "state_output": 2 * state_size * tokens,
        }

    def flops_dense_ffn(self, tokens: int = 1) -> Dict[str, int]:
        c, h = self.c, self.h
        mid = c.intermediate_size
        return {
            "gate_proj": self._linear_flops(h, mid, tokens),
            "up_proj": self._linear_flops(h, mid, tokens),
            "silu_mul": 2 * mid * tokens,
            "down_proj": self._linear_flops(mid, h, tokens),
        }

    def flops_moe(self, tokens: int = 1, activated_only: bool = True) -> Dict[str, int]:
        c, h = self.c, self.h
        n_exp = c.num_experts_per_tok if activated_only else c.num_experts
        mid = c.moe_intermediate_size
        shared_mid = c.shared_expert_intermediate_size

        gate_up = self._linear_flops(h, mid, tokens) * 2
        silu_mul = 2 * mid * tokens
        down = self._linear_flops(mid, h, tokens)
        expert_total = (gate_up + silu_mul + down) * n_exp

        shared_total = 0
        if shared_mid > 0:
            s_gate_up = self._linear_flops(h, shared_mid, tokens) * 2
            s_silu = 2 * shared_mid * tokens
            s_down = self._linear_flops(shared_mid, h, tokens)
            shared_total = s_gate_up + s_silu + s_down

        return {
            "router": self._linear_flops(h, c.num_experts, tokens),
            "shared_expert_gate": self._linear_flops(h, 1, tokens) if shared_mid > 0 else 0,
            f"experts(x{n_exp})": expert_total,
            "shared_expert": shared_total,
        }

    def flops_mlp(self, tokens: int = 1, activated_only: bool = True) -> Dict[str, int]:
        if self.c.is_moe:
            return self.flops_moe(tokens=tokens, activated_only=activated_only)
        return self.flops_dense_ffn(tokens=tokens)

    def flops_lm_head(self, tokens: int = 1) -> Dict[str, int]:
        return {"lm_head": self._linear_flops(self.h, self.c.vocab_size, tokens)}

    def profile_prefill(self, tokens: int) -> Dict[str, Dict[str, int]]:
        c = self.c
        full_proj = self.flops_full_attn_proj(tokens)
        full_score = self.flops_full_attn_score(tokens, tokens)
        linear_proj = self.flops_linear_attn_proj(tokens)
        linear_conv = self.flops_linear_attn_conv(tokens)
        linear_rec = self.flops_linear_attn_recurrent(tokens)
        mlp = self.flops_mlp(tokens=tokens, activated_only=True)
        mlp_key = "moe_mlp" if c.is_moe else "dense_ffn"
        return {
            "full_attn_proj": {k: v * c.num_full_attn_layers for k, v in full_proj.items()},
            "full_attn_score": {k: v * c.num_full_attn_layers for k, v in full_score.items()},
            "linear_attn_proj": {k: v * c.num_linear_attn_layers for k, v in linear_proj.items()},
            "linear_attn_conv": {k: v * c.num_linear_attn_layers for k, v in linear_conv.items()},
            "linear_attn_recurrent": {k: v * c.num_linear_attn_layers for k, v in linear_rec.items()},
            mlp_key: {k: v * c.num_hidden_layers for k, v in mlp.items()},
            "lm_head": self.flops_lm_head(tokens),
        }

    def profile_decode(self, context_len: int) -> Dict[str, Dict[str, int]]:
        c = self.c
        full_proj = self.flops_full_attn_proj(1)
        full_score = self.flops_full_attn_score(1, context_len)
        linear_proj = self.flops_linear_attn_proj(1)
        linear_conv = self.flops_linear_attn_conv(1)
        linear_rec = self.flops_linear_attn_recurrent(1)
        mlp = self.flops_mlp(tokens=1, activated_only=True)
        mlp_key = "moe_mlp" if c.is_moe else "dense_ffn"
        return {
            "full_attn_proj": {k: v * c.num_full_attn_layers for k, v in full_proj.items()},
            "full_attn_score": {k: v * c.num_full_attn_layers for k, v in full_score.items()},
            "linear_attn_proj": {k: v * c.num_linear_attn_layers for k, v in linear_proj.items()},
            "linear_attn_conv": {k: v * c.num_linear_attn_layers for k, v in linear_conv.items()},
            "linear_attn_recurrent": {k: v * c.num_linear_attn_layers for k, v in linear_rec.items()},
            mlp_key: {k: v * c.num_hidden_layers for k, v in mlp.items()},
            "lm_head": self.flops_lm_head(1),
        }

    def total_prefill_flops(self, tokens: int, batch_size: int = 1) -> float:
        prof = self.profile_prefill(tokens)
        return sum(sum(v.values()) for v in prof.values()) * batch_size

    def total_decode_flops(self, context_len: int, batch_size: int = 1) -> float:
        prof = self.profile_decode(context_len)
        return sum(sum(v.values()) for v in prof.values()) * batch_size

    def kv_cache_per_token_per_layer(self) -> float:
        c = self.c
        elements = 2 * c.num_key_value_heads * c.head_dim
        return elements * self.bpa

    def linear_state_per_layer(self) -> Dict[str, float]:
        c = self.c
        conv_elements = c.lin_conv_dim * c.linear_conv_kernel_dim
        recur_elements = c.linear_num_value_heads * c.linear_key_head_dim * c.linear_value_head_dim
        return {
            "conv_cache": conv_elements * self.bpa,
            "recurrent_state": recur_elements * self.bpa,
            "total": (conv_elements + recur_elements) * self.bpa,
        }

    def total_cache_memory(self, context_len: int, batch_size: int = 1) -> Dict[str, float]:
        c = self.c
        kv_bytes = self.kv_cache_per_token_per_layer() * context_len * c.num_full_attn_layers
        lin_bytes = self.linear_state_per_layer()["total"] * c.num_linear_attn_layers
        return {
            "full_attn_kv_cache": kv_bytes * batch_size,
            "linear_attn_state": lin_bytes * batch_size,
            "total": (kv_bytes + lin_bytes) * batch_size,
        }

    def decode_weight_bytes(self) -> Dict[str, float]:
        c = self.c
        full_attn = sum(self.params_full_attn().values()) * c.num_full_attn_layers
        linear_attn = sum(self.params_linear_attn().values()) * c.num_linear_attn_layers
        layer_norm = sum(self.params_layer_norms().values()) * c.num_hidden_layers
        mlp_active = sum(self.params_mlp(activated_only=True).values()) * c.num_hidden_layers
        lm_head = sum(self.params_lm_head().values())
        final_norm = sum(self.params_final_norm().values())
        mlp_name = "moe_mlp_weights(activated)" if c.is_moe else "dense_ffn_weights"
        return {
            "full_attn_weights": full_attn * self.bpw,
            "linear_attn_weights": linear_attn * self.bpw,
            "layer_norm_weights": layer_norm * self.bpw,
            mlp_name: mlp_active * self.bpw,
            "lm_head": lm_head * self.bpw,
            "final_norm": final_norm * self.bpw,
        }

    def decode_kv_io_bytes(self, context_len: int, batch_size: int = 1) -> Dict[str, float]:
        c = self.c
        kv_read = self.kv_cache_per_token_per_layer() * context_len * c.num_full_attn_layers
        kv_write = self.kv_cache_per_token_per_layer() * c.num_full_attn_layers
        lin_state = self.linear_state_per_layer()["total"] * c.num_linear_attn_layers
        return {
            "kv_cache_read": kv_read * batch_size,
            "kv_cache_write": kv_write * batch_size,
            "linear_state_rw": lin_state * 2 * batch_size,
        }

    def prefill_weight_bytes(self, _: int) -> Dict[str, float]:
        return self.decode_weight_bytes()

    def prefill_kv_io_bytes(self, input_tokens: int, batch_size: int = 1) -> Dict[str, float]:
        c = self.c
        kv_write = self.kv_cache_per_token_per_layer() * input_tokens * c.num_full_attn_layers
        lin_state_write = self.linear_state_per_layer()["total"] * c.num_linear_attn_layers
        out = {
            "kv_cache_write": kv_write * batch_size,
            "linear_state_write": lin_state_write * batch_size,
        }
        if self.include_activation_io:
            activation_io = 2 * self.h * self.bpa * input_tokens * c.num_hidden_layers
            attn_scratch = c.num_attention_heads * input_tokens * input_tokens * self.bpa * c.num_full_attn_layers
            out["activation_io"] = activation_io * batch_size
            out["attn_scratch"] = attn_scratch * batch_size
        return out

    @staticmethod
    def roofline_latency_ms(flops: float, nbytes: float, chip_tops: float, chip_bw_gbs: float) -> Dict[str, float]:
        compute_ms = flops / (chip_tops * 1e12) * 1e3
        bw_ms = nbytes / (chip_bw_gbs * 1e9) * 1e3
        latency_ms = max(compute_ms, bw_ms)
        intensity = flops / nbytes if nbytes > 0 else 0.0
        return {
            "compute_ms": compute_ms,
            "bw_ms": bw_ms,
            "latency_ms": latency_ms,
            "intensity": intensity,
            "bound": "compute" if compute_ms >= bw_ms else "memory",
        }

    def prefill_metrics(self, prefill_tokens: int, chip_tops: float, chip_bw_gbs: float, batch_size: int = 1) -> Dict[str, float]:
        flops = self.total_prefill_flops(prefill_tokens, batch_size=batch_size)
        nbytes = sum(self.prefill_weight_bytes(prefill_tokens).values()) + sum(
            self.prefill_kv_io_bytes(prefill_tokens, batch_size=batch_size).values()
        )
        roof = self.roofline_latency_ms(flops, nbytes, chip_tops, chip_bw_gbs)
        return {
            "tokens": float(prefill_tokens),
            "flops": float(flops),
            "bytes": float(nbytes),
            "compute_ms": float(roof["compute_ms"]),
            "bw_ms": float(roof["bw_ms"]),
            "lat_ms": float(roof["latency_ms"]),
            "latency_ms": float(roof["latency_ms"]),
            "tps": float(prefill_tokens * batch_size / (roof["latency_ms"] / 1e3)),
            "intensity": float(roof["intensity"]),
            "bound": roof["bound"],
        }

    def decode_metrics(self, context_len: int, chip_tops: float, chip_bw_gbs: float, batch_size: int = 1) -> Dict[str, float]:
        flops = self.total_decode_flops(context_len, batch_size=batch_size)
        nbytes = sum(self.decode_weight_bytes().values()) + sum(
            self.decode_kv_io_bytes(context_len, batch_size=batch_size).values()
        )
        roof = self.roofline_latency_ms(flops, nbytes, chip_tops, chip_bw_gbs)
        return {
            "context": float(context_len),
            "flops": float(flops),
            "bytes": float(nbytes),
            "compute_ms": float(roof["compute_ms"]),
            "bw_ms": float(roof["bw_ms"]),
            "lat_ms": float(roof["latency_ms"]),
            "latency_ms": float(roof["latency_ms"]),
            "tps": float(batch_size / (roof["latency_ms"] / 1e3)),
            "intensity": float(roof["intensity"]),
            "bound": roof["bound"],
        }

    def prefill_perf(self, prefill_tokens: int, chip_tops: float, chip_bw_gbs: float, batch_size: int = 1) -> Dict[str, float]:
        return self.prefill_metrics(prefill_tokens, chip_tops, chip_bw_gbs, batch_size=batch_size)

    def decode_perf(self, context_len: int, chip_tops: float, chip_bw_gbs: float, batch_size: int = 1) -> Dict[str, float]:
        return self.decode_metrics(context_len, chip_tops, chip_bw_gbs, batch_size=batch_size)

    def decode_tps_memory_only(self, context_len: int, chip_bw_gbs: float, batch_size: int = 1) -> float:
        nbytes = sum(self.decode_weight_bytes().values()) + sum(
            self.decode_kv_io_bytes(context_len, batch_size=batch_size).values()
        )
        if nbytes <= 0:
            return 0.0
        return batch_size * chip_bw_gbs * 1e9 / nbytes

    def sweep_summary(
        self,
        context_lens: Sequence[int],
        prefill_offset: int,
        prefill_compute_tops: Sequence[float],
        prefill_bw_gbs: float,
        decode_bw_gbs: Sequence[float],
        batch_size: int = 1,
    ) -> List[Dict[str, Any]]:
        rows: List[Dict[str, Any]] = []
        total_weight_bytes = self.total_params() * self.bpw
        active_weight_bytes = self.activated_params() * self.bpw

        for ctx in context_lens:
            prefill_tokens = max(1, ctx - prefill_offset)
            prefill_metrics = self.prefill_metrics(
                prefill_tokens,
                chip_tops=float(prefill_compute_tops[0]),
                chip_bw_gbs=float(prefill_bw_gbs),
                batch_size=batch_size,
            )
            cache = self.total_cache_memory(ctx, batch_size=batch_size)
            kv_cache_gb = float(cache["full_attn_kv_cache"] / 1e9)
            linear_state_gb = float(cache["linear_attn_state"] / 1e9)
            cache_total_gb = float(cache["total"] / 1e9)

            row: Dict[str, Any] = {
                "context_len": int(ctx),
                "prefill_tokens": int(prefill_tokens),
                "prefill_flops": float(prefill_metrics["flops"]),
                "prefill_flops_t": float(prefill_metrics["flops"] / 1e12),
                "kv_cache_gb": kv_cache_gb,
                "linear_state_gb": linear_state_gb,
                "cache_total_gb": cache_total_gb,
                "all_plus_cache_gb": float((total_weight_bytes + cache["total"]) / 1e9),
                "active_plus_cache_gb": float((active_weight_bytes + cache["total"]) / 1e9),
                "prefill_perf": {},
                "decode_tps_by_bw": {},
            }

            for tops in prefill_compute_tops:
                perf = self.prefill_perf(prefill_tokens, float(tops), float(prefill_bw_gbs), batch_size=batch_size)
                row["prefill_perf"][f"{tops:g}"] = {
                    "ttft_s": perf["latency_ms"] / 1e3,
                    "prefill_tps": perf["tps"],
                    "bound": perf["bound"],
                }

            for bw in decode_bw_gbs:
                row["decode_tps_by_bw"][f"{bw:g}"] = self.decode_tps_memory_only(
                    context_len=ctx,
                    chip_bw_gbs=float(bw),
                    batch_size=batch_size,
                )

            rows.append(row)
        return rows


def print_summary_table(rows: Sequence[Dict[str, Any]], prefill_compute_tops: Sequence[float], decode_bw_gbs: Sequence[float]) -> None:
    headers = [
        "Context",
        "Prefill FLOPs(T)",
        "KV(GB)",
        "LinearState(GB)",
        "All+Cache(GB)",
        "Active+Cache(GB)",
    ]
    for tops in prefill_compute_tops:
        headers.append(f"TTFT@{tops:g}T(s)")
        headers.append(f"PF-TPS@{tops:g}T")
    for bw in decode_bw_gbs:
        headers.append(f"Dec-TPS@{bw:g}GB/s")

    table_rows: List[List[str]] = []
    for row in rows:
        line = [
            fmt_context(int(row["context_len"])),
            f"{row['prefill_flops_t']:.2f}",
            f"{row['kv_cache_gb']:.2f}",
            f"{row['linear_state_gb']:.2f}",
            f"{row['all_plus_cache_gb']:.2f}",
            f"{row['active_plus_cache_gb']:.2f}",
        ]
        for tops in prefill_compute_tops:
            perf = row["prefill_perf"][f"{tops:g}"]
            line.append(f"{perf['ttft_s']:.3f}")
            line.append(f"{perf['prefill_tps']:.1f}")
        for bw in decode_bw_gbs:
            line.append(f"{row['decode_tps_by_bw'][f'{bw:g}']:.1f}")
        table_rows.append(line)

    widths = []
    for idx, header in enumerate(headers):
        w = len(header)
        for line in table_rows:
            w = max(w, len(line[idx]))
        widths.append(w)

    sep = "-+-".join("-" * w for w in widths)
    print(" | ".join(h.ljust(widths[i]) for i, h in enumerate(headers)))
    print(sep)
    for line in table_rows:
        print(" | ".join(v.rjust(widths[i]) for i, v in enumerate(line)))


def build_util_rows(args: argparse.Namespace, profiler: QwenProfiler) -> Dict[str, Any]:
    if args.compute_utilization <= 0:
        raise ValueError("--compute-utilization must be > 0")
    if args.bandwidth_utilization <= 0:
        raise ValueError("--bandwidth-utilization must be > 0")

    effective_compute_tops = float(args.compute) * float(args.compute_utilization)
    effective_bandwidth_gbs = float(args.bandwidth) * float(args.bandwidth_utilization)
    total_weight_bytes = profiler.total_params() * profiler.bpw

    rows: List[Dict[str, Any]] = []
    for ctx in args.context_lens:
        prefill_tokens = max(1, ctx - args.prefill_offset)
        prefill = profiler.prefill_metrics(
            prefill_tokens,
            effective_compute_tops,
            effective_bandwidth_gbs,
            batch_size=args.batch_size,
        )
        decode = profiler.decode_metrics(
            ctx,
            effective_compute_tops,
            effective_bandwidth_gbs,
            batch_size=args.batch_size,
        )
        cache = profiler.total_cache_memory(ctx, batch_size=args.batch_size)
        rows.append(
            {
                "context_len": int(ctx),
                "context_k": fmt_context_k(int(ctx)),
                "prefill_tokens": int(prefill_tokens),
                "prefill_flops_t": float(prefill["flops"] / 1e12),
                "kv_cache_gb": float(cache["full_attn_kv_cache"] / 1e9),
                "all_gb": float((total_weight_bytes + cache["total"]) / 1e9),
                "ttft_s": float(prefill["lat_ms"] / 1e3),
                "prefill_tps": float(prefill["tps"]),
                "decode_tps": float(decode["tps"]),
                "prefill_bound": prefill["bound"],
                "decode_bound": decode["bound"],
            }
        )

    return {
        "effective_compute_tops": effective_compute_tops,
        "effective_bandwidth_gbs": effective_bandwidth_gbs,
        "rows": rows,
    }


def print_util_table(rows: Sequence[Dict[str, Any]]) -> None:
    headers = [
        "Context(K)",
        "Prefill(T)",
        "KV-Cache(GB)",
        "All(GB)",
        "TTFT(s)",
        "Prefill-TPS",
        "Decode-TPS",
    ]
    table_rows = [
        [
            str(row["context_k"]),
            f"{row['prefill_flops_t']:.2f}",
            f"{row['kv_cache_gb']:.2f}",
            f"{row['all_gb']:.2f}",
            f"{row['ttft_s']:.3f}",
            f"{row['prefill_tps']:.1f}",
            f"{row['decode_tps']:.1f}",
        ]
        for row in rows
    ]

    widths = []
    for idx, header in enumerate(headers):
        width = len(header)
        for line in table_rows:
            width = max(width, len(line[idx]))
        widths.append(width)

    sep = "-+-".join("-" * width for width in widths)
    print(" | ".join(header.ljust(widths[idx]) for idx, header in enumerate(headers)))
    print(sep)
    for line in table_rows:
        print(" | ".join(value.rjust(widths[idx]) for idx, value in enumerate(line)))


def run_profile_summary(args: argparse.Namespace, model_name: str, profiler: QwenProfiler, cfg: ModelConfig) -> Dict[str, Any]:
    rows = profiler.sweep_summary(
        context_lens=args.context_lens,
        prefill_offset=args.prefill_offset,
        prefill_compute_tops=args.prefill_compute_tops,
        prefill_bw_gbs=args.prefill_bandwidth_gbs,
        decode_bw_gbs=args.decode_bandwidth_gbs,
        batch_size=args.batch_size,
    )

    print("=" * 100)
    print("Qwen Unified Profile (summary)")
    print("=" * 100)
    print(f"model:           {model_name}")
    print(f"config:          {args.config}")
    print(f"hidden/layers:   {cfg.hidden_size} / {cfg.num_hidden_layers}")
    print(f"ffn type:        {'MoE' if cfg.is_moe else 'Dense'}")
    print(f"layer split:     linear={cfg.num_linear_attn_layers}, full={cfg.num_full_attn_layers}")
    print(
        f"dtype(bytes):    weight={args.weight_dtype} ({profiler.bpw:.5g} B), "
        f"activation={args.activation_dtype} ({profiler.bpa:.5g} B)"
    )
    print(
        f"bandwidth model: {'include' if args.include_activation_io else 'exclude'} "
        "activation/scratch IO"
    )
    print(f"total params:    {fmt_num(profiler.total_params())} ({fmt_bytes(profiler.total_params() * profiler.bpw)})")
    print(f"active params:   {fmt_num(profiler.activated_params())} ({fmt_bytes(profiler.activated_params() * profiler.bpw)})")
    print(f"batch size:      {args.batch_size}")
    prefill_expr = "context_len" if args.prefill_offset == 0 else f"context_len - {args.prefill_offset}"
    print(f"prefill tokens:  {prefill_expr}")
    print(f"prefill roofline:{list(args.prefill_compute_tops)} TOPS, BW={args.prefill_bandwidth_gbs:g} GB/s")
    print(f"decode BW only:  {list(args.decode_bandwidth_gbs)} GB/s")
    print()
    print_summary_table(rows, args.prefill_compute_tops, args.decode_bandwidth_gbs)

    payload = {
        "mode": "profile",
        "style": "summary",
        "model_name": model_name,
        "config": args.config,
        "dtype": {
            "weight_dtype": args.weight_dtype,
            "activation_dtype": args.activation_dtype,
            "bytes_per_weight": profiler.bpw,
            "bytes_per_activation": profiler.bpa,
        },
        "include_activation_io": args.include_activation_io,
        "batch_size": args.batch_size,
        "prefill_offset": args.prefill_offset,
        "rows": rows,
    }
    return payload


def run_profile_next(args: argparse.Namespace, model_name: str, profiler: QwenProfiler, cfg: ModelConfig) -> Dict[str, Any]:
    print("=" * 100)
    print("Qwen Unified Profile (next)")
    print("=" * 100)
    print(f"model:           {model_name}")
    print(f"config:          {args.config}")
    print(f"ffn type:        {'MoE' if cfg.is_moe else 'Dense'}")
    print(f"layer split:     linear={cfg.num_linear_attn_layers}, full={cfg.num_full_attn_layers}")
    print(
        f"dtype(bytes):    weight={args.weight_dtype} ({profiler.bpw:.5g} B), "
        f"activation={args.activation_dtype} ({profiler.bpa:.5g} B)"
    )
    print(
        f"bandwidth model: {'include' if args.include_activation_io else 'exclude'} "
        "activation/scratch IO"
    )
    print(f"chip:            {args.compute:g} TOPS, {args.bandwidth:g} GB/s")
    print()

    print("[1] Params")
    total_params = profiler.total_params()
    active_params = profiler.activated_params()
    print(f"  total params:    {fmt_num(total_params)} ({fmt_bytes(total_params * profiler.bpw)})")
    print(f"  active params:   {fmt_num(active_params)} ({fmt_bytes(active_params * profiler.bpw)})")
    print()

    print("[2] Cache/State (batch=1)")
    kv_per_tok = profiler.kv_cache_per_token_per_layer() * cfg.num_full_attn_layers
    lin_state = profiler.linear_state_per_layer()["total"] * cfg.num_linear_attn_layers
    print(f"  KV per token:    {fmt_bytes(kv_per_tok)}")
    print(f"  linear state:    {fmt_bytes(lin_state)}")
    print()

    print("[3] Decode")
    decode_rows: List[Dict[str, float]] = []
    for ctx in args.context_lens:
        metrics = profiler.decode_metrics(ctx, args.compute, args.bandwidth, batch_size=args.batch_size)
        decode_rows.append(metrics)
        print(
            f"  ctx={ctx:>8,d} | flops={fmt_flops(metrics['flops']):>15s} | bytes={fmt_bytes(metrics['bytes']):>12s} | "
            f"lat={fmt_time(metrics['lat_ms']):>10s} | tps={metrics['tps']:>8.1f}"
        )
    print()

    print("[4] Prefill")
    prefill_rows: List[Dict[str, float]] = []
    for tok in args.input_tokens:
        metrics = profiler.prefill_metrics(tok, args.compute, args.bandwidth, batch_size=args.batch_size)
        prefill_rows.append(metrics)
        print(
            f"  tok={tok:>8,d} | flops={fmt_flops(metrics['flops']):>15s} | bytes={fmt_bytes(metrics['bytes']):>12s} | "
            f"ttft={fmt_time(metrics['lat_ms']):>10s} | pf_tps={metrics['tps']:>8.1f}"
        )

    return {
        "mode": "profile",
        "style": "next",
        "model_name": model_name,
        "config": args.config,
        "dtype": {
            "weight_dtype": args.weight_dtype,
            "activation_dtype": args.activation_dtype,
            "bytes_per_weight": profiler.bpw,
            "bytes_per_activation": profiler.bpa,
        },
        "include_activation_io": args.include_activation_io,
        "batch_size": args.batch_size,
        "compute_tops": args.compute,
        "bandwidth_gbs": args.bandwidth,
        "decode": decode_rows,
        "prefill": prefill_rows,
    }


def run_profile_util(args: argparse.Namespace, model_name: str, profiler: QwenProfiler, cfg: ModelConfig) -> Dict[str, Any]:
    util_data = build_util_rows(args, profiler)
    effective_compute_tops = util_data["effective_compute_tops"]
    effective_bandwidth_gbs = util_data["effective_bandwidth_gbs"]
    rows = util_data["rows"]

    print("=" * 100)
    print("Qwen Unified Profile (util)")
    print("=" * 100)
    print(f"model:                {model_name}")
    print(f"config:               {args.config}")
    print(f"ffn type:             {'MoE' if cfg.is_moe else 'Dense'}")
    print(f"layer split:          linear={cfg.num_linear_attn_layers}, full={cfg.num_full_attn_layers}")
    print(f"batch size:           {args.batch_size}")
    print(f"prefill tokens:       context_len - {args.prefill_offset}" if args.prefill_offset else "prefill tokens:       context_len")
    print(f"raw compute/bw:       {args.compute:g} TOPS / {args.bandwidth:g} GB/s")
    print(f"utilization factors:  compute={args.compute_utilization:g}, bandwidth={args.bandwidth_utilization:g}")
    print(f"effective compute/bw: {effective_compute_tops:g} TOPS / {effective_bandwidth_gbs:g} GB/s")
    print()
    print_util_table(rows)

    return {
        "mode": "profile",
        "style": "util",
        "model_name": model_name,
        "config": args.config,
        "dtype": {
            "weight_dtype": args.weight_dtype,
            "activation_dtype": args.activation_dtype,
            "bytes_per_weight": profiler.bpw,
            "bytes_per_activation": profiler.bpa,
        },
        "include_activation_io": args.include_activation_io,
        "batch_size": args.batch_size,
        "compute_tops": args.compute,
        "bandwidth_gbs": args.bandwidth,
        "compute_utilization": args.compute_utilization,
        "bandwidth_utilization": args.bandwidth_utilization,
        "effective_compute_tops": effective_compute_tops,
        "effective_bandwidth_gbs": effective_bandwidth_gbs,
        "prefill_offset": args.prefill_offset,
        "rows": rows,
    }


def excel_summary(args: argparse.Namespace, model_name: str, profiler: QwenProfiler):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError("openpyxl is required for Excel export. Install with: pip install openpyxl") from exc

    rows = profiler.sweep_summary(
        context_lens=args.context_lens,
        prefill_offset=args.prefill_offset,
        prefill_compute_tops=args.prefill_compute_tops,
        prefill_bw_gbs=args.prefill_bandwidth_gbs,
        decode_bw_gbs=args.decode_bandwidth_gbs,
        batch_size=args.batch_size,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = safe_sheet_name(model_name, "Summary")

    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=10, bold=True)
    font_normal = Font(name="Calibri", size=10)
    fill_title = PatternFill("solid", fgColor="2F5597")
    fill_group = PatternFill("solid", fgColor="D9E1F2")
    fill_header = PatternFill("solid", fgColor="EDEDED")
    fill_white = PatternFill("solid", fgColor="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def set_cell(row: int, col: int, value: Any, **styles) -> None:
        cell = ws.cell(row=row, column=col, value=value)
        for k, v in styles.items():
            setattr(cell, k, v)

    total_cols = 6 + 2 * len(args.prefill_compute_tops) + len(args.decode_bandwidth_gbs)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title = f"{model_name} (batch={args.batch_size}, {args.weight_dtype}/{args.activation_dtype})"
    set_cell(1, 1, title, font=font_title, fill=fill_title, alignment=align_center, border=border)
    for c in range(2, total_cols + 1):
        set_cell(1, c, None, fill=fill_title, border=border)

    group_row = 2
    header_row = 3
    col = 1
    set_cell(group_row, col, "", font=font_header, fill=fill_group, alignment=align_center, border=border)
    col += 1
    set_cell(group_row, col, "", font=font_header, fill=fill_group, alignment=align_center, border=border)
    col += 1
    ws.merge_cells(start_row=group_row, start_column=col, end_row=group_row, end_column=col + 3)
    set_cell(group_row, col, "Memory", font=font_header, fill=fill_group, alignment=align_center, border=border)
    for c in range(col + 1, col + 4):
        set_cell(group_row, c, None, fill=fill_group, border=border)
    col += 4

    for tops in args.prefill_compute_tops:
        ws.merge_cells(start_row=group_row, start_column=col, end_row=group_row, end_column=col + 1)
        set_cell(group_row, col, f"{tops:g}T TOPS", font=font_header, fill=fill_group, alignment=align_center, border=border)
        set_cell(group_row, col + 1, None, fill=fill_group, border=border)
        col += 2

    for bw in args.decode_bandwidth_gbs:
        set_cell(group_row, col, f"{bw:g}GB/s BW", font=font_header, fill=fill_group, alignment=align_center, border=border)
        col += 1

    headers = [
        "Context",
        "Prefill FLOPs (T)",
        "KV (GB)",
        "Linear State (GB)",
        "All+Cache (GB)",
        "Active+Cache (GB)",
    ]
    for _ in args.prefill_compute_tops:
        headers.extend(["TTFT (s)", "Prefill-TPS"])
    for _ in args.decode_bandwidth_gbs:
        headers.append("Decode-TPS")

    for i, h in enumerate(headers, start=1):
        set_cell(header_row, i, h, font=font_header, fill=fill_header, alignment=align_center, border=border)

    row_idx = 4
    for row in rows:
        col = 1
        vals = [
            fmt_context(int(row["context_len"])),
            float(row["prefill_flops_t"]),
            float(row["kv_cache_gb"]),
            float(row["linear_state_gb"]),
            float(row["all_plus_cache_gb"]),
            float(row["active_plus_cache_gb"]),
        ]
        for v in vals:
            set_cell(row_idx, col, v, font=font_normal, fill=fill_white, alignment=align_center if col == 1 else align_right, border=border)
            if col != 1:
                ws.cell(row=row_idx, column=col).number_format = "0.00"
            col += 1

        for tops in args.prefill_compute_tops:
            perf = row["prefill_perf"][f"{tops:g}"]
            set_cell(row_idx, col, float(perf["ttft_s"]), font=font_normal, fill=fill_white, alignment=align_right, border=border)
            ws.cell(row=row_idx, column=col).number_format = "0.000"
            col += 1
            set_cell(row_idx, col, float(perf["prefill_tps"]), font=font_normal, fill=fill_white, alignment=align_right, border=border)
            ws.cell(row=row_idx, column=col).number_format = "0.0"
            col += 1

        for bw in args.decode_bandwidth_gbs:
            set_cell(
                row_idx,
                col,
                float(row["decode_tps_by_bw"][f"{bw:g}"]),
                font=font_normal,
                fill=fill_white,
                alignment=align_right,
                border=border,
            )
            ws.cell(row=row_idx, column=col).number_format = "0.0"
            col += 1
        row_idx += 1

    widths = [12, 16, 10, 14, 14, 16] + [10, 13] * len(args.prefill_compute_tops) + [12] * len(args.decode_bandwidth_gbs)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A4"

    output_path = make_output_path(model_name, "summary", "profile", args.output)
    wb.save(output_path)
    print(f"excel saved: {output_path}")


def excel_util(args: argparse.Namespace, model_name: str, profiler: QwenProfiler):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError("openpyxl is required for Excel export. Install with: pip install openpyxl") from exc

    util_data = build_util_rows(args, profiler)
    rows = util_data["rows"]

    wb = Workbook()
    ws = wb.active
    ws.title = safe_sheet_name(model_name, "Util")

    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")
    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=10, bold=True)
    font_value = Font(name="Calibri", size=12, bold=True)
    font_cell = Font(name="Calibri", size=10)

    fill_title = PatternFill("solid", fgColor="2F5EA8")
    fill_info = PatternFill("solid", fgColor="FFF2CC")
    fill_group = PatternFill("solid", fgColor="D9E2F3")
    fill_util = PatternFill("solid", fgColor="E2F0D9")
    fill_header = PatternFill("solid", fgColor="EDEDED")
    fill_body = PatternFill("solid", fgColor="FFFFFF")

    def style(cell, *, font=None, fill=None, alignment=None, number_format: Optional[str] = None) -> None:
        cell.border = border
        if font is not None:
            cell.font = font
        if fill is not None:
            cell.fill = fill
        if alignment is not None:
            cell.alignment = alignment
        if number_format is not None:
            cell.number_format = number_format

    def merged_value(start: str, end: str, value: Any, *, font, fill, alignment) -> None:
        ws.merge_cells(f"{start}:{end}")
        cell = ws[start]
        cell.value = value
        style(cell, font=font, fill=fill, alignment=alignment)
        start_col = ws[start].column
        end_col = ws[end].column
        start_row = ws[start].row
        end_row = ws[end].row
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                style(ws.cell(row=row, column=col), fill=fill, alignment=alignment)

    precision_title = f"{brief_precision_label(args.weight_dtype, 'w')}, {brief_precision_label(args.activation_dtype, 'a')}"

    ws.merge_cells("A1:H1")
    ws["A1"] = f"{model_name} ({precision_title})"
    style(ws["A1"], font=font_title, fill=fill_title, alignment=align_center)

    ws["A2"] = "算力（T-FLOPS）："
    ws["B2"] = float(args.compute)
    ws["A3"] = "带宽（GB/s）："
    ws["B3"] = float(args.bandwidth)
    for ref in ("A2", "A3"):
        style(ws[ref], font=font_header, fill=fill_info, alignment=align_center)
    for ref in ("B2", "B3"):
        style(ws[ref], font=font_value, fill=fill_info, alignment=align_center, number_format="0.##")

    merged_value("C2", "C3", "算力需求\n(T-FLOPS)", font=font_header, fill=fill_group, alignment=align_center)
    merged_value("D2", "E3", "显存占用\n(GB)", font=font_header, fill=fill_group, alignment=align_center)
    merged_value("F2", "G2", "算力利用率", font=font_header, fill=fill_util, alignment=align_center)
    ws["H2"] = "带宽利用率"
    style(ws["H2"], font=font_header, fill=fill_util, alignment=align_center)
    merged_value("F3", "G3", float(args.compute_utilization), font=font_value, fill=fill_util, alignment=align_center)
    ws["H3"] = float(args.bandwidth_utilization)
    style(ws["H3"], font=font_value, fill=fill_util, alignment=align_center, number_format="0.###")
    ws["F3"].number_format = "0.###"

    merged_value("A4", "B4", "Context Length  (K)", font=font_header, fill=fill_header, alignment=align_center)
    headers = {
        "C4": "Prefill",
        "D4": "KV-Cache",
        "E4": "All",
        "F4": "TTFT (s)",
        "G4": "Prefill-TPS",
        "H4": "Decode-TPS",
    }
    for ref, value in headers.items():
        ws[ref] = value
        style(ws[ref], font=font_header, fill=fill_header, alignment=align_center)

    row_idx = 5
    for row in rows:
        merged_value(f"A{row_idx}", f"B{row_idx}", row["context_k"], font=font_cell, fill=fill_body, alignment=align_center)

        data_cells = [
            ("C", row["prefill_flops_t"], "0.00"),
            ("D", row["kv_cache_gb"], "0.00"),
            ("E", row["all_gb"], "0.00"),
            ("F", row["ttft_s"], "0.000"),
            ("G", row["prefill_tps"], "0.0"),
            ("H", row["decode_tps"], "0.0"),
        ]
        for col, value, fmt in data_cells:
            cell = ws[f"{col}{row_idx}"]
            cell.value = float(value)
            style(cell, font=font_cell, fill=fill_body, alignment=align_center if col == "C" else align_right, number_format=fmt)
        row_idx += 1

    for row in range(1, row_idx):
        ws.row_dimensions[row].height = 32
    for col, width in {
        "A": 15,
        "B": 9,
        "C": 16,
        "D": 13,
        "E": 12,
        "F": 13,
        "G": 15,
        "H": 14,
    }.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"

    output_path = make_output_path(model_name, "util", "profile", args.output)
    wb.save(output_path)
    print(f"excel saved: {output_path}")


def excel_next(args: argparse.Namespace, model_name: str, profiler: QwenProfiler, cfg: ModelConfig):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise ModuleNotFoundError("openpyxl is required for Excel export. Install with: pip install openpyxl") from exc

    wb = Workbook()

    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    font_h = Font(name="Calibri", size=11, bold=True)
    font_n = Font(name="Calibri", size=10)
    fill_h = PatternFill("solid", fgColor="D9E1F2")
    align_c = Alignment(horizontal="center", vertical="center")
    align_r = Alignment(horizontal="right", vertical="center")

    def style_row(ws, row: int, cols: int) -> None:
        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = font_h
            cell.fill = fill_h
            cell.alignment = align_c
            cell.border = border

    def style_cell(ws, row: int, col: int, value: Any, right: bool = False) -> None:
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = font_n
        cell.alignment = align_r if right else align_c
        cell.border = border

    # Sheet 1: Config/summary
    ws0 = wb.active
    ws0.title = safe_sheet_name(model_name, "参数")
    for i, w in enumerate([28, 20, 20, 42], start=1):
        ws0.column_dimensions[get_column_letter(i)].width = w
    ws0["A1"] = f"{model_name} Profile Summary"
    ws0["A1"].font = Font(name="Calibri", size=14, bold=True)

    rows_info = [
        ("Config", args.config),
        ("FFN", "MoE" if cfg.is_moe else "Dense"),
        ("Hidden", cfg.hidden_size),
        ("Layers", cfg.num_hidden_layers),
        ("Linear layers", cfg.num_linear_attn_layers),
        ("Full layers", cfg.num_full_attn_layers),
        ("Weight dtype", f"{args.weight_dtype} ({profiler.bpw:.5g} B)"),
        ("Activation dtype", f"{args.activation_dtype} ({profiler.bpa:.5g} B)"),
        ("Include activation IO", args.include_activation_io),
        ("Total params", profiler.total_params()),
        ("Active params", profiler.activated_params()),
        ("Total weight bytes", profiler.total_params() * profiler.bpw),
        ("Active weight bytes", profiler.activated_params() * profiler.bpw),
    ]
    for i, (k, v) in enumerate(rows_info, start=3):
        style_cell(ws0, i, 1, k)
        style_cell(ws0, i, 2, v, right=isinstance(v, (int, float)))
        if isinstance(v, (int, float)):
            ws0.cell(row=i, column=2).number_format = "#,##0.00"

    # Sheet 2: Decode
    ws1 = wb.create_sheet(safe_sheet_name(model_name, "Decode"))
    heads = ["Context", "FLOPs", "Bytes", "Compute(ms)", "BW(ms)", "Latency(ms)", "TPS"]
    ws1.append(heads)
    style_row(ws1, 1, len(heads))
    for ctx in args.context_lens:
        metrics = profiler.decode_metrics(ctx, args.compute, args.bandwidth, batch_size=args.batch_size)
        ws1.append([
            ctx,
            metrics["flops"],
            metrics["bytes"],
            metrics["compute_ms"],
            metrics["bw_ms"],
            metrics["lat_ms"],
            metrics["tps"],
        ])

    # Sheet 3: Prefill
    ws2 = wb.create_sheet(safe_sheet_name(model_name, "Prefill"))
    heads = ["Tokens", "FLOPs", "Bytes", "Compute(ms)", "BW(ms)", "TTFT(ms)", "TPS"]
    ws2.append(heads)
    style_row(ws2, 1, len(heads))
    for tok in args.input_tokens:
        metrics = profiler.prefill_metrics(tok, args.compute, args.bandwidth, batch_size=args.batch_size)
        ws2.append([
            tok,
            metrics["flops"],
            metrics["bytes"],
            metrics["compute_ms"],
            metrics["bw_ms"],
            metrics["lat_ms"],
            metrics["tps"],
        ])

    # Sheet 4: Breakdown
    ws3 = wb.create_sheet(safe_sheet_name(model_name, "Breakdown"))
    heads = ["Phase", "Category", "Op", "FLOPs"]
    ws3.append(heads)
    style_row(ws3, 1, len(heads))

    ctx0 = args.context_lens[0]
    tok0 = args.input_tokens[0]
    p_dec = profiler.profile_decode(ctx0)
    p_pre = profiler.profile_prefill(tok0)
    for cat, ops in p_dec.items():
        for op, fl in ops.items():
            ws3.append(["decode", cat, op, fl])
    for cat, ops in p_pre.items():
        for op, fl in ops.items():
            ws3.append(["prefill", cat, op, fl])

    for ws in [ws1, ws2, ws3]:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = font_n
                c.border = border
                c.alignment = align_r
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 16

    output_path = make_output_path(model_name, "next", "profile", args.output)
    wb.save(output_path)
    print(f"excel saved: {output_path}")


def build_profiler_from_args(args: argparse.Namespace) -> tuple[str, ModelConfig, QwenProfiler]:
    model_name = infer_model_name(args.config)
    cfg = ModelConfig.from_json(args.config)
    bpw = resolve_dtype_bytes(args.weight_dtype, args.weight_bytes)
    bpa = resolve_dtype_bytes(args.activation_dtype, args.activation_bytes)
    profiler = QwenProfiler(
        config=cfg,
        bytes_per_weight=bpw,
        bytes_per_activation=bpa,
        include_activation_io=args.include_activation_io,
    )
    return model_name, cfg, profiler


def add_common_model_args(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--config", type=str, default="weights/Qwen3.5-27B/config.json")
    parser.add_argument("--weight-dtype", type=str, default="ssfpw4")
    parser.add_argument("--activation-dtype", type=str, default="int8")
    parser.add_argument("--weight-bytes", type=float, default=None)
    parser.add_argument("--activation-bytes", type=float, default=None)
    parser.add_argument("--batch-size", type=int, default=1)
    parser.add_argument(
        "--context-lens",
        type=int,
        nargs="+",
        default=[2048, 8192, 16384, 32768, 65536, 131072, 262144, 524288, 1048576, 2097152, 4194304],
    )
    parser.add_argument(
        "--include-activation-io",
        action="store_true",
        help="Include activation/scratch I/O in prefill BW model (default: off)",
    )


def add_summary_args(parser: argparse.ArgumentParser) -> None:
    parser.add_argument(
        "--prefill-offset",
        type=int,
        default=0,
        help="summary mode uses prefill_tokens = max(1, context_len - prefill_offset)",
    )
    parser.add_argument("-A", "--prefill-compute-tops", type=float, nargs="+", default=[200.0, 500.0])
    parser.add_argument("-B", "--prefill-bandwidth-gbs", type=float, default=1000.0)
    parser.add_argument("--decode-bandwidth-gbs", type=float, nargs="+", default=[250.0, 1000.0])


def add_next_args(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--compute", type=float, default=200.0)
    parser.add_argument("--bandwidth", type=float, default=1000.0)
    parser.add_argument(
        "--input-tokens",
        type=int,
        nargs="+",
        default=[2048, 8192, 16384, 32768, 65536, 131072, 262144, 524288, 1048576, 2097152, 4194304],
    )


def add_util_args(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--compute-utilization", type=float, default=1.0)
    parser.add_argument("--bandwidth-utilization", type=float, default=1.0)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="All-in-one Qwen profile + excel")
    sub = parser.add_subparsers(dest="command", required=True)

    p_profile = sub.add_parser("profile", help="Print profile report")
    p_profile.add_argument("--style", choices=["summary", "next", "util"], default="summary")
    add_common_model_args(p_profile)
    add_summary_args(p_profile)
    add_next_args(p_profile)
    add_util_args(p_profile)
    p_profile.add_argument("--json-out", type=str, default=None)
    p_profile.add_argument("--log-file", type=str, default=None)
    p_profile.add_argument("--no-log", action="store_true")

    p_excel = sub.add_parser("excel", help="Generate excel report")
    p_excel.add_argument("--style", choices=["summary", "next", "util"], default="summary")
    add_common_model_args(p_excel)
    add_summary_args(p_excel)
    add_next_args(p_excel)
    add_util_args(p_excel)
    p_excel.add_argument("-o", "--output", type=str, default=None)

    return parser.parse_args()


def run_profile(args: argparse.Namespace) -> None:
    model_name, cfg, profiler = build_profiler_from_args(args)

    def _run() -> Dict[str, Any]:
        if args.style == "summary":
            payload = run_profile_summary(args, model_name, profiler, cfg)
        elif args.style == "next":
            payload = run_profile_next(args, model_name, profiler, cfg)
        else:
            payload = run_profile_util(args, model_name, profiler, cfg)

        if args.json_out:
            out_dir = os.path.dirname(args.json_out)
            if out_dir:
                os.makedirs(out_dir, exist_ok=True)
            with open(args.json_out, "w", encoding="utf-8") as f:
                json.dump(payload, f, indent=2, ensure_ascii=False)
            print(f"json saved: {args.json_out}")

        return payload

    if args.no_log:
        _run()
        return

    model_name = infer_model_name(args.config)
    log_path = args.log_file or make_output_path(model_name, args.style, "log", None)
    with open(log_path, "w", encoding="utf-8") as f:
        with redirect_stdout(f):
            _run()
    print(f"profile log saved: {log_path}", file=sys.stderr)


def run_excel(args: argparse.Namespace) -> None:
    model_name, cfg, profiler = build_profiler_from_args(args)
    if args.style == "summary":
        excel_summary(args, model_name, profiler)
    elif args.style == "next":
        excel_next(args, model_name, profiler, cfg)
    else:
        excel_util(args, model_name, profiler)


def main() -> None:
    args = parse_args()
    if args.command == "profile":
        run_profile(args)
    elif args.command == "excel":
        run_excel(args)
    else:
        raise ValueError(f"unsupported command: {args.command}")


if __name__ == "__main__":
    main()
