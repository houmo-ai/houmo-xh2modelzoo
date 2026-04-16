import argparse
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import torch
import torch.nn as nn
from accelerate import init_empty_weights
from loguru import logger
from transformers import AutoConfig, AutoModelForCausalLM, PreTrainedModel


@dataclass
class _ProfileModelConfig:
    hidden_size: int = 2048
    num_hidden_layers: int = 48
    head_dim: int = 256
    vocab_size: int = 151936
    num_attention_heads: int = 16
    num_key_value_heads: int = 2
    linear_key_head_dim: int = 128
    linear_num_key_heads: int = 16
    linear_value_head_dim: int = 128
    linear_num_value_heads: int = 32
    linear_conv_kernel_dim: int = 4
    num_experts: int = 0
    num_experts_per_tok: int = 0
    moe_intermediate_size: int = 0
    shared_expert_intermediate_size: int = 0
    intermediate_size: int = 5120
    decoder_sparse_step: int = 1
    full_attention_interval: int = 4
    layer_types_cfg: list[str] | None = None

    @property
    def is_moe(self) -> bool:
        return self.num_experts > 0 and self.num_experts_per_tok > 0

    @property
    def num_full_attn_layers(self) -> int:
        if self.layer_types_cfg and len(self.layer_types_cfg) == self.num_hidden_layers:
            return sum(1 for layer_type in self.layer_types_cfg if layer_type == "full_attention")
        return sum(1 for index in range(self.num_hidden_layers) if (index + 1) % self.full_attention_interval == 0)

    @property
    def num_linear_attn_layers(self) -> int:
        return self.num_hidden_layers - self.num_full_attn_layers

    @property
    def lin_key_dim(self) -> int:
        return self.linear_num_key_heads * self.linear_key_head_dim

    @property
    def lin_value_dim(self) -> int:
        return self.linear_num_value_heads * self.linear_value_head_dim

    @property
    def lin_conv_dim(self) -> int:
        return self.lin_key_dim + self.lin_key_dim + self.lin_value_dim

    @property
    def full_q_out_dim(self) -> int:
        return self.num_attention_heads * self.head_dim

    @property
    def full_qz_out_dim(self) -> int:
        return self.num_attention_heads * self.head_dim * 2

    @property
    def full_kv_out_dim(self) -> int:
        return self.num_key_value_heads * self.head_dim


@dataclass(frozen=True)
class _HardwareConfig:
    compute: float
    bandwidth: float
    compute_utilization: float
    bandwidth_utilization: float
    c2c_multiplier: float

    @property
    def effective_compute_tops(self) -> float:
        return self.compute * self.compute_utilization

    @property
    def effective_bandwidth_gbs(self) -> float:
        return self.bandwidth * self.bandwidth_utilization * self.c2c_multiplier


@dataclass(frozen=True)
class _ContextSummaryRow:
    context_length_k: int | float
    prefill_tflops: float
    kv_cache_gib: float
    weight_gib: float
    weight_plus_kv_gib: float
    ttft_seconds: float | None
    prefill_tps: float
    decode_tps: float


@dataclass(frozen=True)
class _OutputTargets:
    output_dir: Path
    log_file: Path | None
    excel_file: Path | None


class _ProfileCalculator:
    def __init__(self, config: _ProfileModelConfig, bytes_per_param: int = 2):
        self.c = config
        self.bpp = bytes_per_param
        self.h = config.hidden_size

    @staticmethod
    def _linear_flops(in_dim: int, out_dim: int, tokens: int = 1) -> int:
        return 2 * in_dim * out_dim * tokens

    def params_lm_head(self) -> dict[str, int]:
        return {"lm_head": self.h * self.c.vocab_size}

    def params_final_norm(self) -> dict[str, int]:
        return {"final_norm": self.h}

    def params_full_attn(self) -> dict[str, int]:
        c, hidden_size = self.c, self.h
        return {
            "Q_proj": hidden_size * c.full_q_out_dim,
            "Z_proj": hidden_size * c.full_q_out_dim,
            "K_proj": hidden_size * c.full_kv_out_dim,
            "V_proj": hidden_size * c.full_kv_out_dim,
            "O_proj": c.full_q_out_dim * hidden_size,
            "q_norm": c.head_dim,
            "k_norm": c.head_dim,
        }

    def params_linear_attn(self) -> dict[str, int]:
        c, hidden_size = self.c, self.h
        return {
            "Q_proj": hidden_size * c.lin_key_dim,
            "K_proj": hidden_size * c.lin_key_dim,
            "V_proj": hidden_size * c.lin_value_dim,
            "Z_proj": hidden_size * c.lin_value_dim,
            "A_proj": hidden_size * c.linear_num_value_heads,
            "G_proj": hidden_size * c.linear_num_value_heads,
            "O_proj": c.lin_value_dim * hidden_size,
            "conv1d": c.lin_conv_dim * c.linear_conv_kernel_dim,
            "dt_bias": c.linear_num_value_heads,
            "A_log": c.linear_num_value_heads,
            "gated_norm": c.linear_value_head_dim,
        }

    def params_layer_norms(self) -> dict[str, int]:
        return {
            "input_layernorm": self.h,
            "post_attn_layernorm": self.h,
        }

    def params_dense_ffn(self) -> dict[str, int]:
        c, hidden_size = self.c, self.h
        return {
            "gate_proj": hidden_size * c.intermediate_size,
            "up_proj": hidden_size * c.intermediate_size,
            "down_proj": c.intermediate_size * hidden_size,
        }

    def params_moe(self, activated_only: bool = False) -> dict[str, int]:
        c, hidden_size = self.c, self.h
        num_experts = c.num_experts_per_tok if activated_only else c.num_experts
        expert_params = (
            hidden_size * c.moe_intermediate_size
            + hidden_size * c.moe_intermediate_size
            + c.moe_intermediate_size * hidden_size
        )
        shared_params = (
            hidden_size * c.shared_expert_intermediate_size
            + hidden_size * c.shared_expert_intermediate_size
            + c.shared_expert_intermediate_size * hidden_size
        )
        return {
            "router": hidden_size * c.num_experts,
            "shared_expert_gate": hidden_size,
            f"experts(x{num_experts})": expert_params * num_experts,
            "shared_expert": shared_params,
        }

    def params_mlp(self, activated_only: bool = False) -> dict[str, int]:
        if self.c.is_moe:
            return self.params_moe(activated_only=activated_only)
        return self.params_dense_ffn()

    def flops_full_attn_proj(self, tokens: int = 1) -> dict[str, int]:
        c, hidden_size = self.c, self.h
        return {
            "Q_proj": self._linear_flops(hidden_size, c.full_q_out_dim, tokens),
            "Z_proj": self._linear_flops(hidden_size, c.full_q_out_dim, tokens),
            "K_proj": self._linear_flops(hidden_size, c.full_kv_out_dim, tokens),
            "V_proj": self._linear_flops(hidden_size, c.full_kv_out_dim, tokens),
            "O_proj": self._linear_flops(c.full_q_out_dim, hidden_size, tokens),
        }

    def flops_full_attn_score(self, tokens: int, context_len: int) -> dict[str, int]:
        c = self.c
        num_heads = c.num_attention_heads
        head_dim = c.head_dim
        return {
            "QK^T": 2 * num_heads * head_dim * tokens * context_len,
            "softmax": 5 * num_heads * tokens * context_len,
            "score@V": 2 * num_heads * head_dim * tokens * context_len,
        }

    def flops_linear_attn_proj(self, tokens: int = 1) -> dict[str, int]:
        c, hidden_size = self.c, self.h
        return {
            "Q_proj": self._linear_flops(hidden_size, c.lin_key_dim, tokens),
            "K_proj": self._linear_flops(hidden_size, c.lin_key_dim, tokens),
            "V_proj": self._linear_flops(hidden_size, c.lin_value_dim, tokens),
            "Z_proj": self._linear_flops(hidden_size, c.lin_value_dim, tokens),
            "A_proj": self._linear_flops(hidden_size, c.linear_num_value_heads, tokens),
            "G_proj": self._linear_flops(hidden_size, c.linear_num_value_heads, tokens),
            "O_proj": self._linear_flops(c.lin_value_dim, hidden_size, tokens),
        }

    def flops_linear_attn_conv(self, tokens: int = 1) -> dict[str, int]:
        c = self.c
        return {
            "conv1d": 2 * c.lin_conv_dim * c.linear_conv_kernel_dim * tokens,
            "silu": 4 * c.lin_conv_dim * tokens,
        }

    def flops_linear_attn_recurrent(self, tokens: int = 1) -> dict[str, int]:
        c = self.c
        state_size = c.linear_num_value_heads * c.linear_key_head_dim * c.linear_value_head_dim
        return {
            "state_gate": state_size * tokens,
            "state_update": 2 * state_size * tokens,
            "state_output": 2 * state_size * tokens,
        }

    def flops_dense_ffn(self, tokens: int = 1) -> dict[str, int]:
        c, hidden_size = self.c, self.h
        intermediate_size = c.intermediate_size
        return {
            "gate_proj": self._linear_flops(hidden_size, intermediate_size, tokens),
            "up_proj": self._linear_flops(hidden_size, intermediate_size, tokens),
            "silu_mul": 2 * intermediate_size * tokens,
            "down_proj": self._linear_flops(intermediate_size, hidden_size, tokens),
        }

    def flops_moe(self, tokens: int = 1, activated_only: bool = True) -> dict[str, int]:
        c, hidden_size = self.c, self.h
        num_experts = c.num_experts_per_tok if activated_only else c.num_experts
        gate_up = self._linear_flops(hidden_size, c.moe_intermediate_size, tokens) * 2
        silu_mul = 2 * c.moe_intermediate_size * tokens
        down = self._linear_flops(c.moe_intermediate_size, hidden_size, tokens)
        expert_total = (gate_up + silu_mul + down) * num_experts

        shared_gate_up = self._linear_flops(hidden_size, c.shared_expert_intermediate_size, tokens) * 2
        shared_silu = 2 * c.shared_expert_intermediate_size * tokens
        shared_down = self._linear_flops(c.shared_expert_intermediate_size, hidden_size, tokens)

        return {
            "router": self._linear_flops(hidden_size, c.num_experts, tokens),
            "shared_expert_gate": self._linear_flops(hidden_size, 1, tokens),
            f"experts(x{num_experts})": expert_total,
            "shared_expert": shared_gate_up + shared_silu + shared_down,
        }

    def flops_mlp(self, tokens: int = 1, activated_only: bool = True) -> dict[str, int]:
        if self.c.is_moe:
            return self.flops_moe(tokens=tokens, activated_only=activated_only)
        return self.flops_dense_ffn(tokens=tokens)

    def flops_lm_head(self, tokens: int = 1) -> dict[str, int]:
        return {"lm_head": self._linear_flops(self.h, self.c.vocab_size, tokens)}

    def kv_cache_per_token_per_layer(self) -> int:
        c = self.c
        elements = 2 * c.num_key_value_heads * c.head_dim
        return elements * self.bpp

    def linear_state_per_layer(self) -> dict[str, int]:
        c = self.c
        conv_elements = c.lin_conv_dim * c.linear_conv_kernel_dim
        recur_elements = c.linear_num_value_heads * c.linear_key_head_dim * c.linear_value_head_dim
        return {
            "conv_cache": conv_elements * self.bpp,
            "recurrent_state": recur_elements * self.bpp,
            "total": (conv_elements + recur_elements) * self.bpp,
        }

    def decode_weight_bytes(self) -> dict[str, int]:
        c = self.c
        full_attn_bytes = sum(self.params_full_attn().values()) * self.bpp * c.num_full_attn_layers
        full_norm_bytes = sum(self.params_layer_norms().values()) * self.bpp * c.num_full_attn_layers
        linear_attn_bytes = sum(self.params_linear_attn().values()) * self.bpp * c.num_linear_attn_layers
        linear_norm_bytes = sum(self.params_layer_norms().values()) * self.bpp * c.num_linear_attn_layers
        mlp_bytes = sum(self.params_mlp(activated_only=True).values()) * self.bpp * c.num_hidden_layers
        mlp_label = "moe_weights(activated)" if c.is_moe else "dense_ffn_weights"
        head_bytes = sum(self.params_lm_head().values()) * self.bpp
        final_norm_bytes = sum(self.params_final_norm().values()) * self.bpp
        return {
            "full_attn_weights": full_attn_bytes + full_norm_bytes,
            "linear_attn_weights": linear_attn_bytes + linear_norm_bytes,
            mlp_label: mlp_bytes,
            "lm_head": head_bytes,
            "final_norm": final_norm_bytes,
        }

    def decode_kv_io_bytes(self, context_len: int, batch_size: int = 1) -> dict[str, int]:
        c = self.c
        kv_read = self.kv_cache_per_token_per_layer() * context_len * c.num_full_attn_layers
        kv_write = self.kv_cache_per_token_per_layer() * c.num_full_attn_layers
        linear_state = self.linear_state_per_layer()["total"] * c.num_linear_attn_layers
        linear_rw = linear_state * 2
        return {
            "kv_cache_read": kv_read * batch_size,
            "kv_cache_write": kv_write * batch_size,
            "linear_state_rw": linear_rw * batch_size,
        }

    def prefill_weight_bytes(self, input_tokens: int) -> dict[str, int]:
        del input_tokens
        return self.decode_weight_bytes()

    def prefill_kv_io_bytes(self, input_tokens: int, batch_size: int = 1) -> dict[str, int]:
        c = self.c
        kv_write = self.kv_cache_per_token_per_layer() * input_tokens * c.num_full_attn_layers
        linear_state = self.linear_state_per_layer()["total"] * c.num_linear_attn_layers
        activation_io = 2 * self.h * self.bpp * input_tokens * c.num_hidden_layers
        attn_scratch = c.num_attention_heads * input_tokens * input_tokens * self.bpp * c.num_full_attn_layers
        return {
            "kv_cache_write": kv_write * batch_size,
            "linear_state_write": linear_state * batch_size,
            "activation_io": activation_io * batch_size,
            "attn_scratch": attn_scratch * batch_size,
        }

    def profile_decode(self, context_len: int) -> dict[str, dict[str, int]]:
        c = self.c
        full_proj = self.flops_full_attn_proj(tokens=1)
        full_score = self.flops_full_attn_score(tokens=1, context_len=context_len)
        linear_proj = self.flops_linear_attn_proj(tokens=1)
        linear_conv = self.flops_linear_attn_conv(tokens=1)
        linear_recurrent = self.flops_linear_attn_recurrent(tokens=1)
        mlp = self.flops_mlp(tokens=1, activated_only=True)
        mlp_key = "moe_mlp" if c.is_moe else "dense_ffn"
        return {
            "full_attn_proj": {name: value * c.num_full_attn_layers for name, value in full_proj.items()},
            "full_attn_score": {name: value * c.num_full_attn_layers for name, value in full_score.items()},
            "linear_attn_proj": {name: value * c.num_linear_attn_layers for name, value in linear_proj.items()},
            "linear_attn_conv": {name: value * c.num_linear_attn_layers for name, value in linear_conv.items()},
            "linear_attn_recurrent": {
                name: value * c.num_linear_attn_layers for name, value in linear_recurrent.items()
            },
            mlp_key: {name: value * c.num_hidden_layers for name, value in mlp.items()},
            "lm_head": self.flops_lm_head(tokens=1),
        }

    def profile_prefill(self, input_tokens: int) -> dict[str, dict[str, int]]:
        c = self.c
        full_proj = self.flops_full_attn_proj(tokens=input_tokens)
        full_score = self.flops_full_attn_score(tokens=input_tokens, context_len=input_tokens)
        linear_proj = self.flops_linear_attn_proj(tokens=input_tokens)
        linear_conv = self.flops_linear_attn_conv(tokens=input_tokens)
        linear_recurrent = self.flops_linear_attn_recurrent(tokens=input_tokens)
        mlp = self.flops_mlp(tokens=input_tokens, activated_only=True)
        mlp_key = "moe_mlp" if c.is_moe else "dense_ffn"
        return {
            "full_attn_proj": {name: value * c.num_full_attn_layers for name, value in full_proj.items()},
            "full_attn_score": {name: value * c.num_full_attn_layers for name, value in full_score.items()},
            "linear_attn_proj": {name: value * c.num_linear_attn_layers for name, value in linear_proj.items()},
            "linear_attn_conv": {name: value * c.num_linear_attn_layers for name, value in linear_conv.items()},
            "linear_attn_recurrent": {
                name: value * c.num_linear_attn_layers for name, value in linear_recurrent.items()
            },
            mlp_key: {name: value * c.num_hidden_layers for name, value in mlp.items()},
            "lm_head": self.flops_lm_head(tokens=input_tokens),
        }


def get_empty_hf_model(model_or_config: str, **kwargs) -> Any:
    """仅加载 transformers 模型结构，不初始化权重。"""
    try:
        from transformers.modeling_utils import no_init_weights
    except ImportError:
        no_init_weights = init_empty_weights

    model_or_config_path = Path(model_or_config)
    pretrained_source = str(model_or_config_path) if model_or_config_path.exists() else model_or_config

    config = AutoConfig.from_pretrained(pretrained_source, trust_remote_code=True)
    causal_lm_config = config
    text_config = _get_text_config(config)
    required_causal_lm_attrs = ("vocab_size", "hidden_size", "num_hidden_layers")
    missing_required_attrs = any(
        getattr(config, attr_name, None) is None for attr_name in required_causal_lm_attrs
    )
    if text_config is not config and missing_required_attrs:
        causal_lm_config = text_config
        if not getattr(causal_lm_config, "_name_or_path", None):
            causal_lm_config._name_or_path = getattr(config, "_name_or_path", pretrained_source)

    kwargs.setdefault("trust_remote_code", True)

    with no_init_weights(), init_empty_weights():
        hf_model = AutoModelForCausalLM.from_config(causal_lm_config, **kwargs)

    return hf_model


def _resolve_torch_dtype(dtype_like: Any) -> torch.dtype | None:
    if isinstance(dtype_like, torch.dtype):
        return dtype_like
    if dtype_like is None:
        return None

    dtype_name = str(dtype_like).replace("torch.", "")
    return getattr(torch, dtype_name, None)


def _get_dtype_name(dtype: torch.dtype | None) -> str:
    return str(dtype).replace("torch.", "") if dtype is not None else "unknown"


def _get_dtype_num_bytes(dtype: torch.dtype | None) -> int:
    if dtype is None:
        return 2
    return torch.tensor([], dtype=dtype).element_size()


def _bytes_to_mib(num_bytes: int) -> float:
    return num_bytes / 1024**2


def _bytes_to_gib(num_bytes: int) -> float:
    return num_bytes / 1024**3


def _get_model_dtype(hf_model: PreTrainedModel) -> torch.dtype:
    for parameter in hf_model.parameters():
        return parameter.dtype

    config = getattr(hf_model, "config", None)
    torch_dtype = _resolve_torch_dtype(getattr(config, "torch_dtype", None))
    return torch_dtype or torch.float16


def _get_text_config(config: Any) -> Any:
    for attr_name in ("text_config", "language_config"):
        sub_config = getattr(config, attr_name, None)
        if sub_config is not None:
            return sub_config
    return config


def _resolve_context_length(config: Any, context_length: int | None) -> int:
    if context_length is not None:
        return int(context_length)

    candidates = (
        getattr(config, "max_position_embeddings", None),
        getattr(config, "max_sequence_length", None),
        getattr(config, "context_length", None),
        getattr(config, "model_max_length", None),
        getattr(config, "n_positions", None),
        getattr(config, "seq_length", None),
        getattr(config, "sliding_window", None),
    )
    for candidate in candidates:
        if isinstance(candidate, int) and candidate > 0:
            return candidate
    return 0


def _get_nested_attr(obj: Any, attr_chain: tuple[str, ...]) -> Any:
    current = obj
    for attr_name in attr_chain:
        current = getattr(current, attr_name, None)
        if current is None:
            return None
    return current


def _find_decoder_layers(hf_model: PreTrainedModel) -> list[Any]:
    candidate_paths = (
        ("model",),
        ("language_model",),
        ("text_model",),
        ("transformer",),
        ("model", "model"),
        ("language_model", "model"),
        ("text_model", "model"),
    )
    visited = set()

    for attr_chain in candidate_paths:
        module = _get_nested_attr(hf_model, attr_chain)
        if module is None or id(module) in visited:
            continue
        visited.add(id(module))
        layers = getattr(module, "layers", None)
        if isinstance(layers, (nn.ModuleList, list, tuple)) and len(layers) > 0:
            return list(layers)

    for module in hf_model.modules():
        layers = getattr(module, "layers", None)
        if isinstance(layers, (nn.ModuleList, list, tuple)) and len(layers) > 0:
            return list(layers)

    return []


def _get_layer_types(layers: list[Any], config: Any) -> list[str | None]:
    config_layer_types = getattr(config, "layer_types", None)
    if config_layer_types is not None and len(config_layer_types) == len(layers):
        return list(config_layer_types)
    return [getattr(layer, "layer_type", None) for layer in layers]


def _sum_module_parameters(module: nn.Module | None) -> int:
    if module is None:
        return 0
    return sum(parameter.numel() for parameter in module.parameters())


def _get_first_parameter_dtype(module: nn.Module | None) -> torch.dtype | None:
    if module is None:
        return None
    for parameter in module.parameters():
        return parameter.dtype
    return None


def _infer_num_routed_experts(experts_module: Any, moe_module: nn.Module) -> int | None:
    if isinstance(experts_module, nn.ModuleList):
        return len(experts_module)

    for owner in (experts_module, moe_module, getattr(moe_module, "config", None)):
        if owner is None:
            continue
        for attr_name in ("num_experts", "num_local_experts", "num_routed_experts"):
            value = getattr(owner, attr_name, None)
            if isinstance(value, int) and value > 0:
                return value

    for parameter in experts_module.parameters():
        if parameter.ndim > 0 and parameter.shape[0] > 1:
            return int(parameter.shape[0])

    return None


def _infer_top_k(moe_module: nn.Module, model_config: Any) -> int | None:
    for owner in (
        moe_module,
        getattr(moe_module, "gate", None),
        getattr(moe_module, "router", None),
        getattr(moe_module, "config", None),
        model_config,
    ):
        if owner is None:
            continue
        for attr_name in ("top_k", "num_experts_per_tok"):
            value = getattr(owner, attr_name, None)
            if isinstance(value, int) and value > 0:
                return value
    return None


def _collect_moe_meta(hf_model: PreTrainedModel, total_parameters: int) -> dict[str, Any] | None:
    model_config = _get_text_config(hf_model.config)
    block_infos = []
    visited_expert_modules = set()
    routed_parameters_total = 0
    active_routed_parameters_total = 0.0
    shared_parameters_total = 0
    estimated = False

    for module in hf_model.modules():
        routed_experts = None
        for attr_name in ("routed_experts", "experts"):
            candidate = getattr(module, attr_name, None)
            if candidate is None:
                continue
            if not isinstance(candidate, nn.Module):
                continue
            if not any(parameter.numel() > 0 for parameter in candidate.parameters()):
                continue
            if not (
                hasattr(module, "gate")
                or hasattr(module, "router")
                or hasattr(module, "top_k")
                or hasattr(module, "num_experts_per_tok")
                or hasattr(module, "shared_expert")
                or hasattr(module, "shared_experts")
            ):
                continue
            routed_experts = candidate
            break

        if routed_experts is None or id(routed_experts) in visited_expert_modules:
            continue

        visited_expert_modules.add(id(routed_experts))

        num_routed_experts = _infer_num_routed_experts(routed_experts, module)
        top_k = _infer_top_k(module, model_config)
        routed_parameters = _sum_module_parameters(routed_experts)
        if routed_parameters == 0:
            continue

        shared_parameters = 0
        for attr_name in ("shared_expert", "shared_experts"):
            shared_parameters += _sum_module_parameters(getattr(module, attr_name, None))

        routed_parameters_total += routed_parameters
        shared_parameters_total += shared_parameters

        active_routed_parameters = None
        block_is_estimated = True
        if num_routed_experts is not None and top_k is not None and num_routed_experts > 0:
            active_experts = min(top_k, num_routed_experts)
            if isinstance(routed_experts, nn.ModuleList) and len(routed_experts) > 0:
                expert_parameter_counts = [_sum_module_parameters(expert_module) for expert_module in routed_experts]
                if len(set(expert_parameter_counts)) == 1:
                    active_routed_parameters = expert_parameter_counts[0] * active_experts
                    block_is_estimated = False
                else:
                    active_routed_parameters = (
                        sum(expert_parameter_counts) * active_experts / len(expert_parameter_counts)
                    )
            else:
                active_routed_parameters = routed_parameters * active_experts / num_routed_experts

        if active_routed_parameters is None:
            estimated = True
            continue

        active_routed_parameters_total += active_routed_parameters
        estimated = estimated or block_is_estimated
        block_infos.append(
            {
                "module_type": type(module).__name__,
                "num_routed_experts": num_routed_experts,
                "num_active_experts_per_token": top_k,
                "routed_expert_parameters": routed_parameters,
                "active_routed_expert_parameters": int(active_routed_parameters),
                "shared_expert_parameters": shared_parameters,
                "active_parameter_count_is_estimated": block_is_estimated,
            }
        )

    if not block_infos:
        return None

    active_num_parameters = total_parameters - routed_parameters_total + int(active_routed_parameters_total)
    return {
        "num_moe_layers": len(block_infos),
        "routed_expert_parameters": routed_parameters_total,
        "shared_expert_parameters": shared_parameters_total,
        "active_routed_expert_parameters": int(active_routed_parameters_total),
        "active_num_parameters": active_num_parameters,
        "active_parameter_count_is_estimated": estimated,
        "blocks": block_infos,
    }


def _collect_kv_cache_meta(
    hf_model: PreTrainedModel,
    batch_size: int,
    context_length: int | None,
) -> dict[str, Any]:
    model_dtype = _get_model_dtype(hf_model)
    model_config = _get_text_config(hf_model.config)
    resolved_context_length = _resolve_context_length(model_config, context_length)
    layers = _find_decoder_layers(hf_model)
    layer_types = _get_layer_types(layers, model_config)

    per_layer = []

    for layer_idx, layer in enumerate(layers):
        layer_type = layer_types[layer_idx] if layer_idx < len(layer_types) else None
        if layer_type == "linear_attention":
            continue

        attn = getattr(layer, "self_attn", None) or getattr(layer, "attn", None)
        if attn is None:
            continue

        num_key_value_heads = getattr(attn, "num_key_value_heads", None) or getattr(
            model_config, "num_key_value_heads", None
        )
        num_attention_heads = getattr(attn, "num_heads", None) or getattr(model_config, "num_attention_heads", None)
        if num_key_value_heads is None:
            num_key_value_heads = num_attention_heads

        head_dim = getattr(attn, "head_dim", None) or getattr(model_config, "head_dim", None)
        if head_dim is None and num_attention_heads:
            hidden_size = getattr(model_config, "hidden_size", None)
            if hidden_size is not None:
                head_dim = hidden_size // num_attention_heads

        if num_key_value_heads is None or head_dim is None:
            continue

        sliding_window = getattr(attn, "sliding_window", None)
        if sliding_window in (None, 0, -1):
            if layer_type == "sliding_attention":
                sliding_window = getattr(model_config, "sliding_window", None)
            elif getattr(model_config, "use_sliding_window", False):
                sliding_window = getattr(model_config, "sliding_window", None)

        cache_length = resolved_context_length
        if isinstance(sliding_window, int) and sliding_window > 0:
            cache_length = sliding_window if cache_length == 0 else min(cache_length, sliding_window)

        cache_dtype = (
            _get_first_parameter_dtype(getattr(attn, "k_proj", None))
            or _get_first_parameter_dtype(getattr(attn, "q_proj", None))
            or model_dtype
        )
        dtype_bytes = _get_dtype_num_bytes(cache_dtype)
        bytes_per_token = 2 * batch_size * int(num_key_value_heads) * int(head_dim) * dtype_bytes
        total_bytes = bytes_per_token * cache_length

        per_layer.append(
            {
                "layer_index": layer_idx,
                "layer_type": layer_type or "attention",
                "cache_shape": [batch_size, int(num_key_value_heads), cache_length, int(head_dim)],
                "cache_dtype": _get_dtype_name(cache_dtype),
                "bytes_per_token": bytes_per_token,
                "total_bytes": total_bytes,
                "total_mib": _bytes_to_mib(total_bytes),
                "sliding_window": sliding_window,
            }
        )

    if not per_layer:
        num_hidden_layers = getattr(model_config, "num_hidden_layers", 0)
        num_key_value_heads = getattr(model_config, "num_key_value_heads", None) or getattr(
            model_config, "num_attention_heads", None
        )
        num_attention_heads = getattr(model_config, "num_attention_heads", None)
        head_dim = getattr(model_config, "head_dim", None)
        if head_dim is None and num_attention_heads:
            hidden_size = getattr(model_config, "hidden_size", None)
            if hidden_size is not None:
                head_dim = hidden_size // num_attention_heads

        if num_hidden_layers and num_key_value_heads and head_dim:
            dtype_bytes = _get_dtype_num_bytes(model_dtype)
            bytes_per_token = 2 * batch_size * int(num_key_value_heads) * int(head_dim) * dtype_bytes
            for layer_idx in range(num_hidden_layers):
                total_bytes = bytes_per_token * resolved_context_length
                per_layer.append(
                    {
                        "layer_index": layer_idx,
                        "layer_type": "attention",
                        "cache_shape": [batch_size, int(num_key_value_heads), resolved_context_length, int(head_dim)],
                        "cache_dtype": _get_dtype_name(model_dtype),
                        "bytes_per_token": bytes_per_token,
                        "total_bytes": total_bytes,
                        "total_mib": _bytes_to_mib(total_bytes),
                        "sliding_window": None,
                    }
                )

    total_bytes = sum(item["total_bytes"] for item in per_layer)
    return {
        "batch_size": batch_size,
        "context_length": resolved_context_length,
        "num_layers": len(per_layer),
        "total_bytes": total_bytes,
        "total_mib": _bytes_to_mib(total_bytes),
        "total_gib": _bytes_to_gib(total_bytes),
        "bytes_per_token": sum(item["bytes_per_token"] for item in per_layer),
        "per_layer": per_layer,
    }


def _collect_linear_attention_cache_meta(hf_model: PreTrainedModel, batch_size: int) -> dict[str, Any] | None:
    model_dtype = _get_model_dtype(hf_model)
    model_config = _get_text_config(hf_model.config)
    layers = _find_decoder_layers(hf_model)
    layer_types = _get_layer_types(layers, model_config)
    per_layer = []

    for layer_idx, layer in enumerate(layers):
        layer_type = layer_types[layer_idx] if layer_idx < len(layer_types) else None
        linear_attn = getattr(layer, "linear_attn", None)
        if linear_attn is None and layer_type != "linear_attention":
            continue
        if linear_attn is None:
            continue

        conv_dim = getattr(linear_attn, "conv_dim", None)
        conv_kernel_size = getattr(linear_attn, "conv_kernel_size", None)
        num_v_heads = getattr(linear_attn, "num_v_heads", None)
        head_k_dim = getattr(linear_attn, "head_k_dim", None)
        head_v_dim = getattr(linear_attn, "head_v_dim", None)

        if None in (conv_dim, conv_kernel_size, num_v_heads, head_k_dim, head_v_dim):
            continue

        cache_dtype = _get_first_parameter_dtype(getattr(linear_attn, "conv1d", None)) or model_dtype
        dtype_bytes = _get_dtype_num_bytes(cache_dtype)
        conv_cache_shape = [batch_size, int(conv_dim), int(conv_kernel_size)]
        recurrent_state_shape = [batch_size, int(num_v_heads), int(head_k_dim), int(head_v_dim)]
        conv_cache_bytes = batch_size * int(conv_dim) * int(conv_kernel_size) * dtype_bytes
        recurrent_state_bytes = batch_size * int(num_v_heads) * int(head_k_dim) * int(head_v_dim) * dtype_bytes
        total_bytes = conv_cache_bytes + recurrent_state_bytes

        per_layer.append(
            {
                "layer_index": layer_idx,
                "conv_cache_shape": conv_cache_shape,
                "recurrent_state_shape": recurrent_state_shape,
                "cache_dtype": _get_dtype_name(cache_dtype),
                "conv_cache_bytes": conv_cache_bytes,
                "recurrent_state_bytes": recurrent_state_bytes,
                "total_bytes": total_bytes,
                "total_mib": _bytes_to_mib(total_bytes),
            }
        )

    if not per_layer:
        return None

    total_bytes = sum(item["total_bytes"] for item in per_layer)
    return {
        "batch_size": batch_size,
        "num_layers": len(per_layer),
        "total_bytes": total_bytes,
        "total_mib": _bytes_to_mib(total_bytes),
        "total_gib": _bytes_to_gib(total_bytes),
        "per_layer": per_layer,
    }


def get_model_meta(
    hf_model: PreTrainedModel,
    batch_size: int = 1,
    context_length: int | None = None,
) -> dict[str, Any]:
    """获取模型元信息，包括参数量、MoE 激活参数与缓存占用。"""
    total_parameters = sum(parameter.numel() for parameter in hf_model.parameters())
    trainable_parameters = sum(parameter.numel() for parameter in hf_model.parameters() if parameter.requires_grad)
    model_dtype = _get_model_dtype(hf_model)
    model_config = _get_text_config(hf_model.config)
    kv_cache_meta = _collect_kv_cache_meta(hf_model, batch_size=batch_size, context_length=context_length)
    linear_attention_cache_meta = _collect_linear_attention_cache_meta(hf_model, batch_size=batch_size)
    moe_meta = _collect_moe_meta(hf_model, total_parameters=total_parameters)

    total_decode_state_bytes = kv_cache_meta["total_bytes"]
    if linear_attention_cache_meta is not None:
        total_decode_state_bytes += linear_attention_cache_meta["total_bytes"]

    meta = {
        "architecture": type(hf_model).__name__,
        "model_type": getattr(model_config, "model_type", type(model_config).__name__),
        "dtype": _get_dtype_name(model_dtype),
        "batch_size": batch_size,
        "context_length": kv_cache_meta["context_length"],
        "num_hidden_layers": getattr(model_config, "num_hidden_layers", len(_find_decoder_layers(hf_model))),
        "num_parameters": total_parameters,
        "num_trainable_parameters": trainable_parameters,
        "is_moe": moe_meta is not None,
        "active_num_parameters": moe_meta["active_num_parameters"] if moe_meta is not None else None,
        "kv_cache": kv_cache_meta,
        "linear_attention_cache": linear_attention_cache_meta,
        "total_decode_state_bytes": total_decode_state_bytes,
        "total_decode_state_mib": _bytes_to_mib(total_decode_state_bytes),
        "total_decode_state_gib": _bytes_to_gib(total_decode_state_bytes),
    }
    if moe_meta is not None:
        meta["moe"] = moe_meta
    return meta


def _format_compact_count(value: int | None) -> str:
    if value is None:
        return "-"

    magnitude = abs(value)
    units = (
        (10**12, "T"),
        (10**9, "B"),
        (10**6, "M"),
        (10**3, "K"),
    )
    for divisor, suffix in units:
        if magnitude >= divisor:
            return f"{value:,} ({value / divisor:.2f}{suffix})"
    return f"{value:,}"


def _format_bytes_human(num_bytes: int | None) -> str:
    if num_bytes is None:
        return "-"

    units = (
        (1024**4, "TiB"),
        (1024**3, "GiB"),
        (1024**2, "MiB"),
        (1024, "KiB"),
    )
    for divisor, suffix in units:
        if abs(num_bytes) >= divisor:
            return f"{num_bytes / divisor:.2f} {suffix}"
    return f"{num_bytes} B"


def _format_shape(shape: list[int] | None) -> str:
    if shape is None:
        return "-"
    return " x ".join(str(item) for item in shape)


def _format_metric(value: float | int | None, unit: str = "") -> str:
    if value is None:
        return "-"

    if isinstance(value, int):
        rendered = f"{value}"
    else:
        rendered = f"{value:.4f}".rstrip("0").rstrip(".")

    return f"{rendered} {unit}".strip()


def _stringify_cell(value: Any) -> str:
    if value is None:
        return "-"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".")
    if isinstance(value, (list, tuple)):
        return _format_shape(list(value))
    return str(value)


def _make_ascii_table(headers: list[str], rows: list[list[Any]]) -> str:
    rendered_headers = [_stringify_cell(header) for header in headers]
    rendered_rows = [[_stringify_cell(cell) for cell in row] for row in rows]
    widths = [len(header) for header in rendered_headers]

    for row in rendered_rows:
        for index, cell in enumerate(row):
            widths[index] = max(widths[index], len(cell))

    border = "+-" + "-+-".join("-" * width for width in widths) + "-+"
    header_line = (
        "| " + " | ".join(header.ljust(width) for header, width in zip(rendered_headers, widths, strict=True)) + " |"
    )
    row_lines = [
        "| " + " | ".join(cell.ljust(width) for cell, width in zip(row, widths, strict=True)) + " |"
        for row in rendered_rows
    ]
    return "\n".join([border, header_line, border, *row_lines, border])


def _make_section(title: str, headers: list[str], rows: list[list[Any]]) -> str:
    return f"{title}\n{_make_ascii_table(headers, rows)}"


def _build_context_length_rows(
    hf_model: PreTrainedModel,
    batch_size: int,
    context_lens: list[int],
) -> list[list[Any]]:
    rows = []
    for context_len in _normalize_context_lens(context_lens):
        meta = get_model_meta(hf_model, batch_size=batch_size, context_length=context_len)
        kv_cache_bytes = meta["kv_cache"]["total_bytes"]
        linear_cache = meta.get("linear_attention_cache")
        linear_cache_bytes = linear_cache["total_bytes"] if linear_cache is not None else 0
        rows.append(
            [
                context_len,
                _format_bytes_human(kv_cache_bytes),
                _format_bytes_human(linear_cache_bytes),
                _format_bytes_human(meta["total_decode_state_bytes"]),
            ]
        )
    return rows


def _build_profiler_model_config(hf_model: PreTrainedModel) -> _ProfileModelConfig:
    model_config = _get_text_config(hf_model.config)
    payload = {}

    for field_name in _ProfileModelConfig.__dataclass_fields__:
        if field_name == "layer_types_cfg":
            continue
        value = getattr(model_config, field_name, None)
        if value is not None:
            payload[field_name] = value

    layer_types = getattr(model_config, "layer_types", None)
    if layer_types is not None:
        payload["layer_types_cfg"] = list(layer_types)

    profiler_config = _ProfileModelConfig(**payload)
    if profiler_config.moe_intermediate_size <= 0:
        profiler_config.moe_intermediate_size = profiler_config.intermediate_size
    return profiler_config


def _calculate_stage_latency_seconds(
    total_flops: float,
    total_bytes: float,
    effective_compute_tops: float,
    effective_bandwidth_gbs: float,
) -> float:
    compute_seconds = float("inf")
    if effective_compute_tops > 0:
        compute_seconds = total_flops / (effective_compute_tops * 1e12)

    bandwidth_seconds = float("inf")
    if effective_bandwidth_gbs > 0:
        bandwidth_seconds = total_bytes / (effective_bandwidth_gbs * 1e9)

    return max(compute_seconds, bandwidth_seconds)


def _format_context_length_k(context_length: int) -> int | float:
    context_length_k = context_length / 1024
    return int(context_length_k) if context_length_k.is_integer() else context_length_k


def _normalize_context_lens(context_lens: list[int] | None) -> list[int]:
    if not context_lens:
        return []
    return sorted({int(context_len) for context_len in context_lens if int(context_len) > 0})


def _sum_profile_flops(profile: dict[str, dict[str, int]], batch_size: int) -> int:
    return sum(sum(operation_values.values()) for operation_values in profile.values()) * batch_size


def _calculate_tokens_per_second(token_count: int, batch_size: int, latency_seconds: float) -> float:
    if latency_seconds in (0, float("inf")):
        return 0.0
    return token_count * batch_size / latency_seconds


def _build_excel_summary_rows(
    hf_model: PreTrainedModel,
    meta: dict[str, Any],
    context_lens: list[int],
    hardware_config: _HardwareConfig,
    bytes_per_param: int,
) -> list[_ContextSummaryRow]:
    profiler = _ProfileCalculator(_build_profiler_model_config(hf_model), bytes_per_param=bytes_per_param)
    batch_size = int(meta["batch_size"])
    decode_weight_bytes_map = profiler.decode_weight_bytes()
    decode_weight_bytes = sum(decode_weight_bytes_map.values())
    weight_gib = _bytes_to_gib(decode_weight_bytes)
    rows = []

    for context_len in _normalize_context_lens(context_lens):
        context_meta = get_model_meta(hf_model, batch_size=batch_size, context_length=context_len)
        kv_cache_gib = _bytes_to_gib(context_meta["total_decode_state_bytes"])

        prefill_profile = profiler.profile_prefill(context_len)
        prefill_flops = _sum_profile_flops(prefill_profile, batch_size)
        prefill_total_bytes = sum(profiler.prefill_weight_bytes(context_len).values()) + sum(
            profiler.prefill_kv_io_bytes(context_len, batch_size).values()
        )
        ttft_seconds = _calculate_stage_latency_seconds(
            prefill_flops,
            prefill_total_bytes,
            hardware_config.effective_compute_tops,
            hardware_config.effective_bandwidth_gbs,
        )

        decode_profile = profiler.profile_decode(context_len)
        decode_flops = _sum_profile_flops(decode_profile, batch_size)
        decode_total_bytes = decode_weight_bytes + sum(profiler.decode_kv_io_bytes(context_len, batch_size).values())
        decode_latency_seconds = _calculate_stage_latency_seconds(
            decode_flops,
            decode_total_bytes,
            hardware_config.effective_compute_tops,
            hardware_config.effective_bandwidth_gbs,
        )

        prefill_tps = _calculate_tokens_per_second(context_len, batch_size, ttft_seconds)
        decode_tps = _calculate_tokens_per_second(1, batch_size, decode_latency_seconds)

        rows.append(
            _ContextSummaryRow(
                context_length_k=_format_context_length_k(context_len),
                prefill_tflops=prefill_flops / 1e12,
                kv_cache_gib=kv_cache_gib,
                weight_gib=weight_gib,
                weight_plus_kv_gib=weight_gib + kv_cache_gib,
                ttft_seconds=None if ttft_seconds == float("inf") else ttft_seconds,
                prefill_tps=prefill_tps,
                decode_tps=decode_tps,
            )
        )

    return rows


def _resolve_model_name(
    model_or_config: str,
    hf_model: PreTrainedModel,
    explicit_model_name: str | None = None,
) -> str:
    candidates = [
        explicit_model_name,
        getattr(getattr(hf_model, "config", None), "_name_or_path", None),
        model_or_config,
    ]

    for candidate in candidates:
        if not candidate:
            continue
        text = str(candidate).rstrip("/")
        path = Path(text)
        if path.suffix == ".json":
            return path.parent.name or path.stem
        if path.name:
            return path.name
        if "/" in text:
            return text.split("/")[-1]
        return text

    return type(hf_model).__name__


def _sanitize_output_stem(text: str) -> str:
    sanitized = "".join(char if char.isalnum() or char in ("-", "_", ".") else "_" for char in text)
    sanitized = sanitized.strip("._")
    return sanitized or "model_profile"


def _resolve_output_dir(output_dir: str | None, model_name: str) -> Path:
    resolved_output_dir = Path(output_dir) if output_dir is not None else Path("work_dirs") / f"{model_name}_profile"
    resolved_output_dir.mkdir(parents=True, exist_ok=True)
    return resolved_output_dir


def _build_default_prefix(output_dir: Path, model_name: str, compute: float, bandwidth: float) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    compute_tag = str(compute).replace(".", "p")
    bandwidth_tag = str(bandwidth).replace(".", "p")
    return output_dir / f"profile_{model_name}_A{compute_tag}_B{bandwidth_tag}_{timestamp}"


def _resolve_artifact_path(path_value: str | None, output_dir: Path, default_path: Path | None) -> Path | None:
    if path_value is None:
        return default_path

    path = Path(path_value)
    if path.is_absolute():
        return path
    return output_dir / path


def _build_hardware_config(args: argparse.Namespace) -> _HardwareConfig:
    return _HardwareConfig(
        compute=args.compute,
        bandwidth=args.bandwidth,
        compute_utilization=args.compute_utilization,
        bandwidth_utilization=args.bandwidth_utilization,
        c2c_multiplier=args.c2c_multiplier,
    )


def _build_output_targets(args: argparse.Namespace, model_name: str) -> _OutputTargets:
    output_dir = _resolve_output_dir(args.output_dir, model_name)

    default_prefix = None
    if (not args.no_log and args.log_file is None) or (not args.no_excel and args.excel_file is None):
        default_prefix = _build_default_prefix(output_dir, model_name, args.compute, args.bandwidth)

    log_file = None
    if not args.no_log:
        log_file = _resolve_artifact_path(
            args.log_file,
            output_dir,
            default_prefix.with_suffix(".log") if default_prefix is not None else None,
        )

    excel_file = None
    if not args.no_excel:
        excel_file = _resolve_artifact_path(
            args.excel_file,
            output_dir,
            default_prefix.with_suffix(".xlsx") if default_prefix is not None else None,
        )

    return _OutputTargets(output_dir=output_dir, log_file=log_file, excel_file=excel_file)


def _emit_meta_text(meta_text: str, log_file: Path | None) -> None:
    _configure_output_logger(log_file)
    logger.opt(raw=True).info(meta_text + "\n")


def export_model_meta_excel(
    excel_path: Path,
    title: str,
    hf_model: PreTrainedModel,
    meta: dict[str, Any],
    context_lens: list[int],
    hardware_config: _HardwareConfig,
    bytes_per_param: int,
) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"

    summary_rows = _build_excel_summary_rows(
        hf_model=hf_model,
        meta=meta,
        context_lens=_normalize_context_lens(context_lens),
        hardware_config=hardware_config,
        bytes_per_param=bytes_per_param,
    )

    thin_side = Side(style="thin", color="000000")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    title_fill = PatternFill("solid", fgColor="2F5597")
    label_fill = PatternFill("solid", fgColor="FFF2CC")
    group_fill = PatternFill("solid", fgColor="E7E6E6")
    value_fill = PatternFill("solid", fgColor="FFF7DA")
    header_fill = PatternFill("solid", fgColor="F2F2F2")

    title_font = Font(color="FFFFFF", bold=True, size=16)
    bold_font = Font(bold=True, size=11)
    header_font = Font(bold=True, size=10)

    def _apply_style(cell_range: str, *, fill=None, font=None, alignment=None, applied_border=None) -> None:
        for row in worksheet[cell_range]:
            for cell in row:
                if fill is not None:
                    cell.fill = fill
                if font is not None:
                    cell.font = font
                if alignment is not None:
                    cell.alignment = alignment
                if applied_border is not None:
                    cell.border = applied_border

    worksheet.merge_cells("A1:I1")
    worksheet["A1"] = title
    _apply_style("A1:I1", fill=title_fill, font=title_font, alignment=center, applied_border=border)

    worksheet["A2"] = "算力 (T-FLOPS)"
    worksheet["B2"] = float(hardware_config.compute)
    worksheet["A3"] = "带宽 (GB/s)"
    worksheet["B3"] = float(hardware_config.bandwidth)
    _apply_style("A2:A3", fill=label_fill, font=bold_font, alignment=center, applied_border=border)
    _apply_style("B2:B3", fill=value_fill, font=bold_font, alignment=center, applied_border=border)

    worksheet.merge_cells("C2:C3")
    worksheet["C2"] = "算力需求\n(T-FLOPS)"
    _apply_style("C2:C3", fill=group_fill, font=bold_font, alignment=center, applied_border=border)

    worksheet.merge_cells("D2:F3")
    worksheet["D2"] = "显存占用\n(GiB)"
    _apply_style("D2:F3", fill=group_fill, font=bold_font, alignment=center, applied_border=border)

    worksheet["G2"] = "算力利用率"
    worksheet["H2"] = "带宽利用率"
    worksheet["I2"] = "C2C倍率"
    worksheet["G3"] = float(hardware_config.compute_utilization)
    worksheet["H3"] = float(hardware_config.bandwidth_utilization)
    worksheet["I3"] = float(hardware_config.c2c_multiplier)
    _apply_style("G2:I2", fill=label_fill, font=bold_font, alignment=center, applied_border=border)
    _apply_style("G3:I3", fill=value_fill, font=bold_font, alignment=center, applied_border=border)

    worksheet.merge_cells("A4:B4")
    worksheet["A4"] = "Context Length (K)"
    worksheet["C4"] = "Prefill"
    worksheet["D4"] = "KV-Cache"
    worksheet["E4"] = "Weight"
    worksheet["F4"] = "Weight+KV"
    worksheet["G4"] = "TTFT (s)"
    worksheet["H4"] = "Prefill-TPS"
    worksheet["I4"] = "Decode-TPS"
    _apply_style("A4:I4", fill=header_fill, font=header_font, alignment=center, applied_border=border)

    for row_index, row in enumerate(summary_rows, start=5):
        worksheet.merge_cells(f"A{row_index}:B{row_index}")
        worksheet[f"A{row_index}"] = row.context_length_k
        worksheet[f"C{row_index}"] = row.prefill_tflops
        worksheet[f"D{row_index}"] = row.kv_cache_gib
        worksheet[f"E{row_index}"] = row.weight_gib
        worksheet[f"F{row_index}"] = row.weight_plus_kv_gib
        worksheet[f"G{row_index}"] = row.ttft_seconds
        worksheet[f"H{row_index}"] = row.prefill_tps
        worksheet[f"I{row_index}"] = row.decode_tps
        _apply_style(f"A{row_index}:I{row_index}", alignment=center, applied_border=border)

    for column_name, width in {
        "A": 14,
        "B": 10,
        "C": 16,
        "D": 12,
        "E": 11,
        "F": 12,
        "G": 12,
        "H": 14,
        "I": 12,
    }.items():
        worksheet.column_dimensions[column_name].width = width

    worksheet.row_dimensions[1].height = 30
    worksheet.row_dimensions[2].height = 28
    worksheet.row_dimensions[3].height = 26
    worksheet.row_dimensions[4].height = 26
    worksheet.freeze_panes = "A5"

    for row_index in range(2, 4):
        worksheet[f"B{row_index}"].number_format = "0.##"
    for coordinate in ("G3", "H3", "I3"):
        worksheet[coordinate].number_format = "0.##"
    for row_index in range(5, 5 + len(summary_rows)):
        worksheet[f"A{row_index}"].number_format = "0.##"
        worksheet[f"C{row_index}"].number_format = "0.00"
        worksheet[f"D{row_index}"].number_format = "0.00"
        worksheet[f"E{row_index}"].number_format = "0.00"
        worksheet[f"F{row_index}"].number_format = "0.00"
        worksheet[f"G{row_index}"].number_format = "0.00"
        worksheet[f"H{row_index}"].number_format = "0.##"
        worksheet[f"I{row_index}"].number_format = "0.##"

    excel_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(excel_path)


def _configure_output_logger(log_file: Path | None = None) -> None:
    logger.remove()
    logger.add(sys.stdout, format="{message}", level="INFO", colorize=False)

    if log_file is not None:
        log_file.parent.mkdir(parents=True, exist_ok=True)
        logger.add(log_file, format="{message}", level="INFO", encoding="utf-8")


def format_model_meta_tables(
    meta: dict[str, Any],
    hf_model: PreTrainedModel | None = None,
    context_lens: list[int] | None = None,
    hardware_config: _HardwareConfig | None = None,
) -> str:
    sections = []

    if hardware_config is not None:
        hardware_rows = [
            ["chip_compute", _format_metric(hardware_config.compute, "TOPS")],
            ["chip_bandwidth", _format_metric(hardware_config.bandwidth, "GB/s")],
            ["compute_utilization", _format_metric(hardware_config.compute_utilization)],
            ["bandwidth_utilization", _format_metric(hardware_config.bandwidth_utilization)],
            ["c2c_multiplier", _format_metric(hardware_config.c2c_multiplier)],
            ["effective_compute", _format_metric(hardware_config.effective_compute_tops, "TOPS")],
            ["effective_bandwidth", _format_metric(hardware_config.effective_bandwidth_gbs, "GB/s")],
        ]
        sections.append(_make_section("硬件配置", ["字段", "值"], hardware_rows))

    summary_rows = [
        ["architecture", meta.get("architecture")],
        ["model_type", meta.get("model_type")],
        ["dtype", meta.get("dtype")],
        ["batch_size", meta.get("batch_size")],
        ["context_length", meta.get("context_length")],
        ["num_hidden_layers", meta.get("num_hidden_layers")],
        ["num_parameters", _format_compact_count(meta.get("num_parameters"))],
        ["num_trainable_parameters", _format_compact_count(meta.get("num_trainable_parameters"))],
        ["is_moe", meta.get("is_moe")],
        ["active_num_parameters", _format_compact_count(meta.get("active_num_parameters"))],
        ["kv_cache_total", _format_bytes_human(meta["kv_cache"]["total_bytes"])],
        [
            "linear_attention_cache_total",
            _format_bytes_human(
                meta["linear_attention_cache"]["total_bytes"] if meta.get("linear_attention_cache") is not None else 0
            ),
        ],
        ["total_decode_state", _format_bytes_human(meta.get("total_decode_state_bytes"))],
    ]
    sections.append(_make_section("Meta 概览", ["字段", "值"], summary_rows))

    moe_meta = meta.get("moe")
    if moe_meta is not None:
        moe_rows = [
            ["num_moe_layers", moe_meta.get("num_moe_layers")],
            ["routed_expert_parameters", _format_compact_count(moe_meta.get("routed_expert_parameters"))],
            [
                "active_routed_expert_parameters",
                _format_compact_count(moe_meta.get("active_routed_expert_parameters")),
            ],
            ["shared_expert_parameters", _format_compact_count(moe_meta.get("shared_expert_parameters"))],
            ["active_num_parameters", _format_compact_count(moe_meta.get("active_num_parameters"))],
            ["estimated", moe_meta.get("active_parameter_count_is_estimated")],
        ]
        sections.append(_make_section("MoE 概览", ["字段", "值"], moe_rows))

        block_rows = []
        for index, block_info in enumerate(moe_meta.get("blocks", []), start=1):
            block_rows.append(
                [
                    index,
                    block_info.get("module_type"),
                    block_info.get("num_routed_experts"),
                    block_info.get("num_active_experts_per_token"),
                    _format_compact_count(block_info.get("routed_expert_parameters")),
                    _format_compact_count(block_info.get("active_routed_expert_parameters")),
                    _format_compact_count(block_info.get("shared_expert_parameters")),
                    block_info.get("active_parameter_count_is_estimated"),
                ]
            )
        if block_rows:
            sections.append(
                _make_section(
                    "MoE Block 明细",
                    [
                        "序号",
                        "module_type",
                        "num_routed_experts",
                        "top_k",
                        "routed_params",
                        "active_params",
                        "shared_params",
                        "estimated",
                    ],
                    block_rows,
                )
            )

    kv_cache_meta = meta["kv_cache"]
    kv_summary_rows = [
        ["batch_size", kv_cache_meta.get("batch_size")],
        ["context_length", kv_cache_meta.get("context_length")],
        ["num_layers", kv_cache_meta.get("num_layers")],
        ["bytes_per_token", _format_bytes_human(kv_cache_meta.get("bytes_per_token"))],
        ["total", _format_bytes_human(kv_cache_meta.get("total_bytes"))],
    ]
    sections.append(_make_section("KV Cache 概览", ["字段", "值"], kv_summary_rows))

    kv_layer_rows = []
    for layer_info in kv_cache_meta.get("per_layer", []):
        kv_layer_rows.append(
            [
                layer_info.get("layer_index"),
                layer_info.get("layer_type"),
                _format_shape(layer_info.get("cache_shape")),
                layer_info.get("cache_dtype"),
                _format_bytes_human(layer_info.get("bytes_per_token")),
                _format_bytes_human(layer_info.get("total_bytes")),
                layer_info.get("sliding_window") if layer_info.get("sliding_window") is not None else "-",
            ]
        )
    if kv_layer_rows:
        sections.append(
            _make_section(
                "KV Cache 分层明细",
                ["layer", "type", "cache_shape", "dtype", "bytes_per_token", "total", "sliding_window"],
                kv_layer_rows,
            )
        )

    linear_attention_cache_meta = meta.get("linear_attention_cache")
    if linear_attention_cache_meta is not None:
        linear_summary_rows = [
            ["batch_size", linear_attention_cache_meta.get("batch_size")],
            ["num_layers", linear_attention_cache_meta.get("num_layers")],
            ["total", _format_bytes_human(linear_attention_cache_meta.get("total_bytes"))],
        ]
        sections.append(_make_section("线性注意力 Cache 概览", ["字段", "值"], linear_summary_rows))

        linear_layer_rows = []
        for layer_info in linear_attention_cache_meta.get("per_layer", []):
            linear_layer_rows.append(
                [
                    layer_info.get("layer_index"),
                    _format_shape(layer_info.get("conv_cache_shape")),
                    _format_shape(layer_info.get("recurrent_state_shape")),
                    layer_info.get("cache_dtype"),
                    _format_bytes_human(layer_info.get("conv_cache_bytes")),
                    _format_bytes_human(layer_info.get("recurrent_state_bytes")),
                    _format_bytes_human(layer_info.get("total_bytes")),
                ]
            )
        if linear_layer_rows:
            sections.append(
                _make_section(
                    "线性注意力 Cache 分层明细",
                    [
                        "layer",
                        "conv_cache_shape",
                        "recurrent_state_shape",
                        "dtype",
                        "conv_cache",
                        "recurrent_state",
                        "total",
                    ],
                    linear_layer_rows,
                )
            )

    if hf_model is not None and context_lens:
        context_rows = _build_context_length_rows(hf_model, meta["batch_size"], context_lens)
        sections.append(
            _make_section(
                "不同 Context Length 下的 Decode State",
                ["context_length", "kv_cache", "linear_cache", "total_decode_state"],
                context_rows,
            )
        )

    return "\n\n".join(sections)


# ============================================================================
# Main
# ============================================================================


def main():
    parser = argparse.ArgumentParser(
        description="Qwen3-Next 模型 Profile 工具",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("--model-name", type=str, help="模型名称")
    parser.add_argument(
        "--config",
        type=str,
        default="weights/Qwen3-Next-80B-A3B-Instruct/config.json",
        help="模型 config.json 路径",
    )
    parser.add_argument(
        "--output-dir",
        type=str,
        help="模型输出目录路径，日志和 Excel 等文件默认都输出到这里",
    )
    parser.add_argument(
        "-A",
        "--compute",
        type=float,
        default=200.0,
        help="芯片算力 (TOPS), 默认 200",
    )
    parser.add_argument(
        "-B",
        "--bandwidth",
        type=float,
        default=272.0,
        help="芯片带宽 (GB/s), 默认 272",
    )
    parser.add_argument(
        "--compute-utilization",
        type=float,
        default=0.5,
        help="算力利用率，默认 0.5",
    )
    parser.add_argument(
        "--bandwidth-utilization",
        type=float,
        default=0.7,
        help="带宽利用率，默认 0.7",
    )
    parser.add_argument(
        "--c2c-multiplier",
        type=float,
        default=1.0,
        help="C2C 倍率，默认 1.0",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=1,
        help="Batch size, 默认 1",
    )
    parser.add_argument(
        "--bytes-per-param",
        type=int,
        default=2,
        help="每参数字节数 (2=bf16, 1=int8, 4=fp32), 默认 2",
    )
    parser.add_argument(
        "--context-lens",
        type=int,
        nargs="+",
        default=[2048, 8192, 16384, 32768, 65536, 131072, 262144, 524288, 1048576, 2097152, 4194304],
        help="Decode 上下文长度列表",
    )
    parser.add_argument(
        "--meta-context-len",
        type=int,
        default=None,
        help="用于 meta 中 cache 占用估算的上下文长度，默认使用模型配置中的长度",
    )
    parser.add_argument(
        "--log-file",
        type=str,
        default=None,
        help="输出日志文件路径 (默认自动生成带时间戳的文件名)",
    )
    parser.add_argument(
        "--excel-file",
        type=str,
        default=None,
        help="Excel 报表输出路径 (默认与日志同前缀，后缀 .xlsx)",
    )
    parser.add_argument(
        "--excel-title",
        type=str,
        default=None,
        help="Excel 报表标题，默认使用模型名",
    )
    parser.add_argument(
        "--no-log",
        action="store_true",
        help="不写日志文件，直接打印到终端 (旧行为)",
    )
    parser.add_argument(
        "--no-excel",
        action="store_true",
        help="不写 Excel 报表",
    )

    args = parser.parse_args()
    context_lens = _normalize_context_lens(args.context_lens)
    hf_model = get_empty_hf_model(args.config)
    resolved_model_name = _resolve_model_name(
        args.config,
        hf_model,
        explicit_model_name=args.model_name,
    )
    hardware_config = _build_hardware_config(args)

    meta = get_model_meta(hf_model, batch_size=args.batch_size, context_length=args.meta_context_len)
    meta_text = format_model_meta_tables(
        meta,
        hf_model=hf_model,
        context_lens=context_lens,
        hardware_config=hardware_config,
    )
    model_name = _sanitize_output_stem(resolved_model_name)
    output_targets = _build_output_targets(args, model_name)

    _emit_meta_text(meta_text, output_targets.log_file)

    if output_targets.excel_file is not None:
        export_model_meta_excel(
            excel_path=output_targets.excel_file,
            title=args.excel_title or resolved_model_name,
            hf_model=hf_model,
            meta=meta,
            context_lens=context_lens,
            hardware_config=hardware_config,
            bytes_per_param=args.bytes_per_param,
        )


if __name__ == "__main__":
    main()
