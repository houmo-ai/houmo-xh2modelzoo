from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import onnx
import openpyxl

try:
    import torch
    from xhquant.api import (
        DeviceType,
        HMONNXInference,
        QuantScheme,
        convert_onnx_to_hmonnx,
        create_quant_config,
    )
    import xhquant.xhonnxruntime.config as xhonnxruntime_config
    try:
        from xhquant.xhonnxruntime.hmonnx_graph_inference import HMONNXGrapInference
    except ImportError:
        HMONNXGrapInference = None
except ModuleNotFoundError as exc:
    raise SystemExit(
        "缺少依赖，请使用 xhquant 环境执行，例如：\n"
        "conda activate xhquant\n"
        "python examples/cv/zte/l3_precision_runner.py"
    ) from exc


REPO_ROOT = Path(__file__).resolve().parents[3]
L3_TEST_ROOT_ENV = "ZTE_L3_TEST_ROOT"
STAGE_EXPORT = "export"
STAGE_INFER = "infer"
STAGE_COMPARE = "compare"
PIPELINE_STAGES = (STAGE_EXPORT, STAGE_INFER, STAGE_COMPARE)


@dataclass(frozen=True)
class CaseConfig:
    name: str
    model_path: Path
    compare_script: Path
    compare_mode: str
    ref_path: Path
    vendor_output_path: Path
    vendor_compare_name: str
    default_tolerance: float | None = None
    preprocess_mode: str | None = None


CASE_NAMES = ("timi", "hxx_lstm")


def _resolve_device(device: str) -> str:
    if device == "auto":
        return "cuda" if torch.cuda.is_available() else "cpu"
    return device


def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def _looks_like_l3_test_root(path: Path) -> bool:
    timi_ok = (
        (path / "timi" / "model_timi.onnx").exists()
        and (path / "timi" / "inference_results.jsonl").exists()
        and (path / "timi" / "64Batch_output_gpu.jsonl").exists()
        and (path / "timi" / "compare_inference_results.py").exists()
    )
    hxx_model_ok = (path / "hxx_lstm" / "model_hxx_v4.onnx").exists() or (path / "hxx_lstm" / "model_hxx_v4_batch1.onnx").exists()
    hxx_ok = (
        hxx_model_ok
        and (path / "hxx_lstm" / "model_io.xlsx").exists()
        and (path / "hxx_lstm" / "gpu_inference_output.xlsx").exists()
        and (path / "hxx_lstm" / "compare_model_outputs.py").exists()
    )
    return timi_ok and hxx_ok


def _normalize_l3_test_root(path: Path) -> Path:
    path = path.expanduser().resolve()
    candidates = [
        path,
        path / "L3精度测试用例",
        path / "L3模型待测试项" / "L3精度测试用例",
    ]
    for candidate in candidates:
        if _looks_like_l3_test_root(candidate):
            return candidate
    return path


def resolve_l3_test_root(explicit_root: Path | None) -> Path:
    candidates: list[Path] = []
    if explicit_root is not None:
        candidates.append(explicit_root)
    env_root = os.environ.get(L3_TEST_ROOT_ENV)
    if env_root:
        candidates.append(Path(env_root))

    candidates.extend(
        [
            REPO_ROOT / "weights/zte/L3模型待测试项/L3精度测试用例",
            REPO_ROOT / "L3模型待测试项/L3精度测试用例",
            REPO_ROOT / "weights/L3模型待测试项/L3精度测试用例",
        ]
    )

    for candidate in candidates:
        normalized = _normalize_l3_test_root(candidate)
        if _looks_like_l3_test_root(normalized):
            return normalized

    for candidate in REPO_ROOT.rglob("L3精度测试用例"):
        if candidate.is_dir() and _looks_like_l3_test_root(candidate):
            return candidate.resolve()

    raise FileNotFoundError(
        "未找到 L3 精度测试数据目录。请通过 `--l3-test-root` 指定 `L3精度测试用例` 目录，"
        f"或设置环境变量 `{L3_TEST_ROOT_ENV}`。"
    )


def build_case_configs(l3_test_root: Path) -> dict[str, CaseConfig]:
    timi_root = l3_test_root / "timi"
    hxx_root = l3_test_root / "hxx_lstm"
    hxx_raw_model = hxx_root / "model_hxx_v4.onnx"
    hxx_batch1_model = hxx_root / "model_hxx_v4_batch1.onnx"

    if hxx_raw_model.exists():
        hxx_model_path = hxx_raw_model
        hxx_preprocess_mode = "fix_hxx_batch1"
    elif hxx_batch1_model.exists():
        hxx_model_path = hxx_batch1_model
        hxx_preprocess_mode = None
    else:
        hxx_model_path = hxx_raw_model
        hxx_preprocess_mode = "fix_hxx_batch1"

    return {
        "timi": CaseConfig(
            name="timi",
            model_path=timi_root / "model_timi.onnx",
            compare_script=timi_root / "compare_inference_results.py",
            compare_mode="jsonl",
            ref_path=timi_root / "inference_results.jsonl",
            vendor_output_path=timi_root / "64Batch_output_gpu.jsonl",
            vendor_compare_name="vendor_gpu_vs_pytorch_ref",
            default_tolerance=1e-5,
            preprocess_mode="fix_timi_batch",
        ),
        "hxx_lstm": CaseConfig(
            name="hxx_lstm",
            model_path=hxx_model_path,
            compare_script=hxx_root / "compare_model_outputs.py",
            compare_mode="excel",
            ref_path=hxx_root / "model_io.xlsx",
            vendor_output_path=hxx_root / "gpu_inference_output.xlsx",
            vendor_compare_name="vendor_gpu_vs_model_io_base",
            preprocess_mode=hxx_preprocess_mode,
        ),
    }


def _load_onnx_input_specs(model_path: Path) -> list[tuple[str, int]]:
    model = onnx.load(str(model_path))
    initializer_names = {initializer.name for initializer in model.graph.initializer}
    specs: list[tuple[str, int]] = []
    for value_info in model.graph.input:
        if value_info.name in initializer_names:
            continue
        specs.append((value_info.name, value_info.type.tensor_type.elem_type))
    return specs


def _cast_numpy_for_runtime(array: np.ndarray, elem_type: int, runtime_float_dtype: torch.dtype) -> np.ndarray:
    if elem_type in (onnx.TensorProto.FLOAT, onnx.TensorProto.FLOAT16, onnx.TensorProto.DOUBLE):
        target_dtype = np.float16 if runtime_float_dtype == torch.float16 else np.float32
        return np.asarray(array, dtype=target_dtype)
    if elem_type == onnx.TensorProto.INT64:
        return np.asarray(array, dtype=np.int32)
    if elem_type in (onnx.TensorProto.INT32, onnx.TensorProto.INT16, onnx.TensorProto.INT8, onnx.TensorProto.UINT8):
        return np.asarray(array)
    if elem_type == onnx.TensorProto.BOOL:
        return np.asarray(array, dtype=bool)
    return np.asarray(array)


def _cast_numpy_for_convert(array: np.ndarray, elem_type: int) -> np.ndarray:
    if elem_type in (onnx.TensorProto.FLOAT, onnx.TensorProto.FLOAT16, onnx.TensorProto.DOUBLE):
        return np.asarray(array, dtype=np.float32)
    if elem_type == onnx.TensorProto.INT64:
        return np.asarray(array, dtype=np.int32)
    if elem_type == onnx.TensorProto.BOOL:
        return np.asarray(array, dtype=bool)
    return np.asarray(array)


def _numpy_to_torch(array: np.ndarray, device: str) -> torch.Tensor:
    tensor = torch.from_numpy(np.ascontiguousarray(array))
    return tensor.to(device)


def _normalize_single_output(output: Any) -> np.ndarray:
    if isinstance(output, (list, tuple)):
        if len(output) != 1:
            raise ValueError(f"期望单输出模型，实际拿到 {len(output)} 个输出")
        output = output[0]
    if isinstance(output, torch.Tensor):
        return output.detach().float().cpu().numpy()
    return np.asarray(output, dtype=np.float32)


def _create_hmonnx_runtime_session(hmonnx_path: Path, device: str) -> tuple[Any, str]:
    runtime_candidates: list[tuple[str, Any]] = []
    if HMONNXGrapInference is not None:
        runtime_candidates.append(("graph", HMONNXGrapInference))
    runtime_candidates.append(("eager", HMONNXInference))

    init_errors: list[str] = []
    for runtime_name, runtime_cls in runtime_candidates:
        try:
            session = runtime_cls(str(hmonnx_path))
            session.to(device)
            session.save_golden = False
            return session, runtime_name
        except Exception as exc:  # pragma: no cover - fallback path depends on runtime internals.
            init_errors.append(f"{runtime_name}: {exc!r}")

    raise RuntimeError(
        f"无法初始化 HMONNX runtime: {hmonnx_path}\n" + "\n".join(init_errors)
    )


def _run_subprocess(cmd: list[str], cwd: Path) -> None:
    env = os.environ.copy()
    env.setdefault("MPLBACKEND", "Agg")
    subprocess.run(cmd, cwd=str(cwd), check=True, env=env)


def _json_default(value: Any) -> Any:
    if isinstance(value, (np.floating, np.integer)):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    raise TypeError(f"Unsupported type for json dump: {type(value)}")


def _load_jsonl(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if line:
                records.append(json.loads(line))
    return records


def _get_runtime_float_dtype(dtype_name: str) -> torch.dtype:
    if dtype_name.lower() in {"fp16", "float16"}:
        return torch.float16
    if dtype_name.lower() in {"fp32", "float32"}:
        return torch.float32
    raise ValueError(f"不支持的 float dtype: {dtype_name}")


def _build_timi_example_inputs(specs: list[tuple[str, int]], ref_path: Path) -> list[np.ndarray]:
    first_record = _load_jsonl(ref_path)[0]
    raw_inputs = first_record["inputs"]
    example_inputs: list[np.ndarray] = []
    for input_name, elem_type in specs:
        example_inputs.append(_cast_numpy_for_convert(np.asarray(raw_inputs[input_name]), elem_type))
    return example_inputs


def _get_timi_batch_size(ref_path: Path) -> int:
    first_record = _load_jsonl(ref_path)[0]
    src = np.asarray(first_record["inputs"]["src"])
    if src.ndim != 3:
        raise ValueError(f"timi src 维度异常: {src.shape}")
    return int(src.shape[0])


def prepare_model_for_case(case_config: CaseConfig, output_root: Path, overwrite: bool) -> Path:
    if case_config.preprocess_mode is None:
        return case_config.model_path

    if not case_config.model_path.exists():
        raise FileNotFoundError(f"{case_config.name} 模型不存在: {case_config.model_path}")

    if case_config.preprocess_mode == "fix_timi_batch":
        batch_size = _get_timi_batch_size(case_config.ref_path)
        prepared_model = output_root / case_config.name / "preprocessed" / f"{case_config.model_path.stem}_batch{batch_size}.onnx"
        if prepared_model.exists() and not overwrite:
            return prepared_model

        prepared_model.parent.mkdir(parents=True, exist_ok=True)
        _run_subprocess(
            [
                sys.executable,
                "-m",
                "onnxsim",
                str(case_config.model_path),
                str(prepared_model),
                "--overwrite-input-shape",
                f"src:{batch_size},512,64",
                f"src_key_padding_mask:{batch_size},512",
            ],
            REPO_ROOT,
        )
        return prepared_model

    if case_config.preprocess_mode == "fix_hxx_batch1":
        prepared_model = output_root / case_config.name / "preprocessed" / f"{case_config.model_path.stem}_batch1.onnx"
        if prepared_model.exists() and not overwrite:
            return prepared_model

        prepared_model.parent.mkdir(parents=True, exist_ok=True)
        _run_subprocess(
            [
                sys.executable,
                "-m",
                "onnxsim",
                str(case_config.model_path),
                str(prepared_model),
                "--input-shape",
                "lstm_input:1,9,13",
            ],
            REPO_ROOT,
        )
        return prepared_model

    raise ValueError(f"未知的 preprocess_mode: {case_config.preprocess_mode}")


def get_expected_output_path(case_config: CaseConfig, output_root: Path, quant_type: str) -> Path:
    case_root = output_root / case_config.name / quant_type / "outputs"
    if case_config.compare_mode == "jsonl":
        return case_root / "hmonnx_output.jsonl"
    return case_root / "hmonnx_output.xlsx"


def get_expected_hmonnx_path(model_path: Path, output_root: Path, case_name: str, quant_type: str) -> Path:
    return output_root / case_name / quant_type / "hmonnx" / f"{model_path.stem}_{quant_type}_XH2a.onnx"


def write_run_manifest(
    output_root: Path,
    l3_test_root: Path,
    case_configs: dict[str, CaseConfig],
    quant_types: list[str],
    overwrite: bool,
) -> Path:
    manifest_path = output_root / "run_manifest.json"
    if manifest_path.exists() and not overwrite:
        return manifest_path

    payload: dict[str, Any] = {
        "generated_at": datetime.now().isoformat(),
        "l3_test_root": str(l3_test_root),
        "stages": list(PIPELINE_STAGES),
        "cases": {},
    }

    for case_name, case_config in case_configs.items():
        prepared_model_path = prepare_model_for_case(case_config, output_root, overwrite=False)
        case_payload = {
            "prepared_model_path": str(prepared_model_path),
            "compare_mode": case_config.compare_mode,
            "expected_output_format": "jsonl" if case_config.compare_mode == "jsonl" else "xlsx",
            "reference_path": str(case_config.ref_path),
            "vendor_output_path": str(case_config.vendor_output_path),
            "compare_script": str(case_config.compare_script),
            "quant_types": {},
        }
        for quant_type in quant_types:
            case_payload["quant_types"][quant_type] = {
                "hmonnx_path": str(get_expected_hmonnx_path(prepared_model_path, output_root, case_name, quant_type)),
                "expected_output_path": str(get_expected_output_path(case_config, output_root, quant_type)),
                "expected_output_note": (
                    "jsonl lines with at least iteration and output fields"
                    if case_config.compare_mode == "jsonl"
                    else "xlsx file with numeric output columns"
                ),
            }
        payload["cases"][case_name] = case_payload

    manifest_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=_json_default) + "\n", encoding="utf-8")
    return manifest_path


def _load_hxx_input_rows(excel_path: Path) -> np.ndarray:
    workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        worksheet = workbook["模型输入"] if "模型输入" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows: list[np.ndarray] = []
        first_row = True
        for row in worksheet.iter_rows(values_only=True):
            if first_row:
                first_row = False
                continue
            values = np.asarray(row, dtype=np.float32)
            if values.size != 117:
                raise ValueError(f"期望 117 个输入值，实际得到 {values.size}")
            rows.append(values.reshape(1, 9, 13))
        return np.asarray(rows, dtype=np.float32)
    finally:
        workbook.close()


def _build_hxx_example_inputs(ref_path: Path) -> list[np.ndarray]:
    inputs = _load_hxx_input_rows(ref_path)
    return [inputs[0]]


def quantize_to_hmonnx(
    case_config: CaseConfig,
    model_path: Path,
    quant_type: str,
    output_root: Path,
    overwrite: bool,
) -> Path:
    hmonnx_path = get_expected_hmonnx_path(model_path, output_root, case_config.name, quant_type)
    if hmonnx_path.exists() and not overwrite:
        return hmonnx_path

    _ensure_parent(hmonnx_path)
    specs = _load_onnx_input_specs(model_path)
    if case_config.name == "timi":
        example_np_inputs = _build_timi_example_inputs(specs, case_config.ref_path)
    elif case_config.name == "hxx_lstm":
        example_np_inputs = _build_hxx_example_inputs(case_config.ref_path)
    else:
        raise KeyError(f"未知 case: {case_config.name}")

    example_inputs = [_numpy_to_torch(array, "cpu") for array in example_np_inputs]
    quant_scheme = QuantScheme(target_device=DeviceType.XH2a, quant_type=quant_type)
    quant_config = create_quant_config(quant_scheme)
    convert_onnx_to_hmonnx(
        str(model_path),
        example_inputs,
        DeviceType.XH2a,
        str(hmonnx_path),
        quant_config=quant_config,
    )
    return hmonnx_path


def run_timi_hmonnx_inference(
    case_config: CaseConfig,
    model_path: Path,
    hmonnx_path: Path,
    output_path: Path,
    device: str,
    runtime_float_dtype: torch.dtype,
) -> dict[str, Any]:
    specs = _load_onnx_input_specs(model_path)
    records = _load_jsonl(case_config.ref_path)
    _ensure_parent(output_path)

    xhonnxruntime_config.disable_progress = True
    xhonnxruntime_config.verbose_progress = False
    session, runtime_backend = _create_hmonnx_runtime_session(hmonnx_path, device)

    latencies_ms: list[float] = []
    with output_path.open("w", encoding="utf-8") as handle:
        for record in records:
            runtime_inputs = []
            for input_name, elem_type in specs:
                array = _cast_numpy_for_runtime(np.asarray(record["inputs"][input_name]), elem_type, runtime_float_dtype)
                runtime_inputs.append(_numpy_to_torch(array, device))

            if device == "cuda":
                torch.cuda.synchronize()
            start_time = time.perf_counter()
            output = session(*runtime_inputs)
            if device == "cuda":
                torch.cuda.synchronize()
            latency_ms = (time.perf_counter() - start_time) * 1000.0
            latencies_ms.append(latency_ms)

            output_array = _normalize_single_output(output)
            handle.write(
                json.dumps(
                    {
                        "iteration": record["iteration"],
                        "latency_ms": latency_ms,
                        "output": output_array.tolist(),
                    },
                    ensure_ascii=False,
                )
                + "\n"
            )

    return {
        "samples": len(records),
        "avg_latency_ms": float(np.mean(latencies_ms)) if latencies_ms else 0.0,
        "runtime_backend": runtime_backend,
        "output_path": str(output_path),
    }


def _write_excel_matrix(matrix: np.ndarray, output_path: Path, sheet_name: str) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    header = [f"y_{index}" for index in range(matrix.shape[1])]
    worksheet.append(header)
    for row in matrix:
        worksheet.append([float(value) for value in row])
    _ensure_parent(output_path)
    workbook.save(output_path)


def run_hxx_hmonnx_inference(
    case_config: CaseConfig,
    model_path: Path,
    hmonnx_path: Path,
    output_path: Path,
    device: str,
    runtime_float_dtype: torch.dtype,
) -> dict[str, Any]:
    inputs = _load_hxx_input_rows(case_config.ref_path)
    runtime_dtype = np.float16 if runtime_float_dtype == torch.float16 else np.float32

    xhonnxruntime_config.disable_progress = True
    xhonnxruntime_config.verbose_progress = False
    session, runtime_backend = _create_hmonnx_runtime_session(hmonnx_path, device)

    latencies_ms: list[float] = []
    outputs: list[np.ndarray] = []
    for sample in inputs:
        tensor = _numpy_to_torch(sample.astype(runtime_dtype, copy=False), device)
        if device == "cuda":
            torch.cuda.synchronize()
        start_time = time.perf_counter()
        output = session(tensor)
        if device == "cuda":
            torch.cuda.synchronize()
        latency_ms = (time.perf_counter() - start_time) * 1000.0
        latencies_ms.append(latency_ms)
        outputs.append(_normalize_single_output(output).reshape(-1))

    stacked_outputs = np.asarray(outputs, dtype=np.float32)
    _write_excel_matrix(stacked_outputs, output_path, sheet_name="模型输出-HMONNX")
    return {
        "samples": int(stacked_outputs.shape[0]),
        "avg_latency_ms": float(np.mean(latencies_ms)) if latencies_ms else 0.0,
        "runtime_backend": runtime_backend,
        "output_path": str(output_path),
    }


def run_inference(
    case_config: CaseConfig,
    model_path: Path,
    hmonnx_path: Path,
    output_root: Path,
    quant_type: str,
    device: str,
    runtime_float_dtype: torch.dtype,
    overwrite: bool,
) -> tuple[Path, dict[str, Any]]:
    output_path = get_expected_output_path(case_config, output_root, quant_type)
    if case_config.compare_mode == "jsonl":
        if output_path.exists() and not overwrite:
            return output_path, {"output_path": str(output_path)}
        stats = run_timi_hmonnx_inference(case_config, model_path, hmonnx_path, output_path, device, runtime_float_dtype)
        return output_path, stats

    if output_path.exists() and not overwrite:
        return output_path, {"output_path": str(output_path)}
    stats = run_hxx_hmonnx_inference(case_config, model_path, hmonnx_path, output_path, device, runtime_float_dtype)
    return output_path, stats


def run_export_stage(
    case_configs: dict[str, CaseConfig],
    cases: list[str],
    quant_types: list[str],
    output_root: Path,
    l3_test_root: Path,
    overwrite: bool,
) -> Path:
    for case_name in cases:
        case_config = case_configs[case_name]
        model_path = prepare_model_for_case(case_config, output_root, overwrite)
        for quant_type in quant_types:
            quantize_to_hmonnx(case_config, model_path, quant_type, output_root, overwrite)
    return write_run_manifest(output_root, l3_test_root, {case: case_configs[case] for case in cases}, quant_types, overwrite)


def run_compare_stage(
    case_configs: dict[str, CaseConfig],
    cases: list[str],
    quant_types: list[str],
    output_root: Path,
    l3_test_root: Path,
    device: str,
    runtime_float_dtype: torch.dtype,
    overwrite: bool,
    skip_vendor_baseline: bool,
) -> tuple[Path, Path]:
    vendor_baselines: dict[str, dict[str, Any]] = {}
    if not skip_vendor_baseline:
        for case_name in cases:
            vendor_baselines[case_name] = run_vendor_baseline(case_configs[case_name], output_root, overwrite)

    quant_results: dict[str, list[dict[str, Any]]] = {case_name: [] for case_name in cases}
    for case_name in cases:
        case_config = case_configs[case_name]
        model_path = prepare_model_for_case(case_config, output_root, overwrite=False)
        for quant_type in quant_types:
            hmonnx_path = get_expected_hmonnx_path(model_path, output_root, case_name, quant_type)
            output_path = get_expected_output_path(case_config, output_root, quant_type)
            if not output_path.exists():
                raise FileNotFoundError(
                    f"缺少推理输出: {output_path}\n"
                    f"请先执行 `--stages {STAGE_INFER}`，或将芯片推理结果写入该路径后再执行 `--stages {STAGE_COMPARE}`。"
                )
            compare_result = compare_outputs(
                case_config=case_config,
                test_output_path=output_path,
                output_root=output_root,
                result_name=quant_type,
                tolerance=case_config.default_tolerance,
            )
            quant_results[case_name].append(
                {
                    "case": case_name,
                    "quant_type": quant_type,
                    "hmonnx_path": str(hmonnx_path),
                    "test_output_path": str(output_path),
                    "infer_stats": {"output_path": str(output_path)},
                    "compare": compare_result,
                }
            )

    summary_json = write_summary_json(output_root, l3_test_root, quant_results, vendor_baselines)
    summary_md = write_summary_markdown(
        output_root=output_root,
        l3_test_root=l3_test_root,
        quant_results=quant_results,
        vendor_baselines=vendor_baselines,
        quant_types=quant_types,
        device=device,
        runtime_float_dtype=runtime_float_dtype,
    )
    return summary_json, summary_md


def compare_timi_outputs(
    case_config: CaseConfig,
    test_output_path: Path,
    output_root: Path,
    result_name: str,
    tolerance: float,
) -> dict[str, Any]:
    report_dir = output_root / case_config.name / result_name / "compare"
    report_dir.mkdir(parents=True, exist_ok=True)
    _run_subprocess(
        [
            sys.executable,
            str(case_config.compare_script),
            "--ref",
            str(case_config.ref_path),
            "--test",
            str(test_output_path),
            "--output",
            str(report_dir),
            "--tolerance",
            str(tolerance),
        ],
        REPO_ROOT,
    )
    with (report_dir / "comparison_report.json").open("r", encoding="utf-8") as handle:
        report = json.load(handle)
    return {
        "report_dir": str(report_dir),
        "summary": report["summary"],
    }


def compare_hxx_outputs(
    case_config: CaseConfig,
    test_output_path: Path,
    output_root: Path,
    result_name: str,
) -> dict[str, Any]:
    report_path = output_root / case_config.name / result_name / "compare" / "report.xlsx"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    _run_subprocess(
        [
            sys.executable,
            str(case_config.compare_script),
            "--base-excel",
            str(case_config.ref_path),
            "--new-excel",
            str(test_output_path),
            "--report-excel",
            str(report_path),
        ],
        REPO_ROOT,
    )
    workbook = openpyxl.load_workbook(report_path, data_only=True, read_only=True)
    try:
        worksheet = workbook["summary"]
        headers = [cell for cell in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        values = [cell for cell in next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True))]
        summary = dict(zip(headers, values))
    finally:
        workbook.close()
    return {
        "report_path": str(report_path),
        "summary": summary,
    }


def compare_outputs(
    case_config: CaseConfig,
    test_output_path: Path,
    output_root: Path,
    result_name: str,
    tolerance: float | None,
) -> dict[str, Any]:
    if case_config.compare_mode == "jsonl":
        return compare_timi_outputs(case_config, test_output_path, output_root, result_name, tolerance or case_config.default_tolerance or 1e-5)
    return compare_hxx_outputs(case_config, test_output_path, output_root, result_name)


def run_vendor_baseline(case_config: CaseConfig, output_root: Path, overwrite: bool) -> dict[str, Any]:
    baseline_dir = output_root / case_config.name / "vendor_baseline"
    summary_path = baseline_dir / "summary.json"
    if summary_path.exists() and not overwrite:
        with summary_path.open("r", encoding="utf-8") as handle:
            summary = json.load(handle)
        # Keep cached baseline metadata aligned with the current display label.
        summary["label"] = case_config.vendor_compare_name
        with summary_path.open("w", encoding="utf-8") as handle:
            json.dump(summary, handle, indent=2, ensure_ascii=False, default=_json_default)
        return summary
    result = compare_outputs(
        case_config=case_config,
        test_output_path=case_config.vendor_output_path,
        output_root=output_root,
        result_name="vendor_baseline",
        tolerance=case_config.default_tolerance,
    )
    summary = {
        "case": case_config.name,
        "label": case_config.vendor_compare_name,
        "test_output_path": str(case_config.vendor_output_path),
        **result,
    }
    baseline_dir.mkdir(parents=True, exist_ok=True)
    with summary_path.open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, indent=2, ensure_ascii=False, default=_json_default)
    return summary


def _format_timi_row(label: str, summary: dict[str, Any]) -> str:
    return (
        f"| {label} | {summary['total_samples']} | {summary['pass_rate'] * 100:.2f}% | "
        f"{summary['numerical_summary']['max_diff_mean']:.6e} | "
        f"{summary['numerical_summary']['max_diff_max']:.6e} | "
        f"{summary['consistency_summary']['mean_accuracy'] * 100:.4f}% |"
    )


def _format_hxx_row(label: str, summary: dict[str, Any]) -> str:
    return f"| {label} | {summary['MAE']:.6e} | {summary['MAPE(%)']:.6f} |"


def write_summary_markdown(
    output_root: Path,
    l3_test_root: Path,
    quant_results: dict[str, list[dict[str, Any]]],
    vendor_baselines: dict[str, dict[str, Any]],
    quant_types: list[str],
    device: str,
    runtime_float_dtype: torch.dtype,
) -> Path:
    summary_path = output_root / "summary.md"
    generated_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    lines: list[str] = []
    lines.append("# ZTE L3 精度测试流程与结果")
    lines.append("")
    lines.append(f"- 生成时间: {generated_time}")
    lines.append(f"- 执行设备: {device}")
    lines.append(f"- HMONNX 浮点输入类型: {runtime_float_dtype}")
    lines.append(f"- 量化设置: {', '.join(quant_types)}")
    lines.append(f"- L3 测试数据目录: {l3_test_root}")
    lines.append(f"- 汇总目录: {output_root}")
    lines.append("")
    lines.append("## 流程")
    lines.append("")
    lines.append("1. 使用 `xhquant` 环境运行 `examples/cv/zte/l3_precision_runner.py`。")
    lines.append("2. `export` 阶段会预处理模型并导出 hmonnx，同时生成 `run_manifest.json`。")
    lines.append("3. `infer` 阶段默认优先使用 `HMONNXGrapInference` 运行推理；若初始化失败，再回退到 `HMONNXInference`。")
    lines.append("4. 如果需要换成芯片推理，可跳过 `infer` 阶段，按 `run_manifest.json` 中的 `expected_output_path` 回填输出。")
    lines.append("5. `compare` 阶段调用对方提供的 compare 脚本，生成 vendor baseline、量化结果报告和 summary。")
    lines.append("")
    lines.append("## 命令")
    lines.append("")
    lines.append("```shell")
    lines.append("conda activate xhquant")
    lines.append(
        "python examples/cv/zte/l3_precision_runner.py "
        "--stages export infer compare "
        "--l3-test-root /path/to/L3精度测试用例 "
        "--cases timi hxx_lstm "
        f"--quant-types {' '.join(quant_types)} "
        "--output-root work_dirs/zte/l3_precision"
    )
    lines.append("```")
    lines.append("")
    lines.append("## 说明")
    lines.append("")
    lines.append("- `timi` 的 compare 脚本要求 `max_diff <= 1e-5` 且标签准确率不低于 99%。")
    lines.append("- 对方 `64Batch_output_gpu.jsonl` 在该阈值下本身也是 `0%` pass rate，因此 `timi` 主要比较 `Avg Max Diff` 和 `Mean Label Acc`。")
    lines.append("- `hxx_lstm` 主要比较 `MAE` 和 `MAPE(%)`，数值越小越好。")
    lines.append(f"- 如果 `L3精度测试用例` 不在仓库默认位置，请显式传 `--l3-test-root`，或设置环境变量 `{L3_TEST_ROOT_ENV}`。")
    lines.append("- 如果要替换成芯片推理，可先执行 `--stages export`，然后把芯片结果写到 `run_manifest.json` 里的 `expected_output_path`，最后执行 `--stages compare`。")
    lines.append("- 运行 backend 默认优先 `HMONNXGrapInference`，因为它基于 graph 执行，通常比 `HMONNXInference` 更快。")
    lines.append("- `Pass Rate`: 满足 compare 脚本通过条件的样本占比；`timi` 的阈值较严，连对方 GPU 结果也可能是 0%。")
    lines.append("- `Avg Max Diff`: 每个样本最大绝对误差的平均值，越小越接近 reference。")
    lines.append("- `Peak Max Diff`: 所有样本里出现过的最大绝对误差峰值，反映最差点位误差。")
    lines.append("- `Mean Label Acc`: 对输出最后一维做 `argmax` 后的平均标签一致率，越高越好。")
    lines.append("- `MAE`: 平均绝对误差，衡量整体数值偏差。")
    lines.append("- `MAPE(%)`: 平均绝对百分比误差，表示相对误差大小，越小越好。")
    lines.append("")

    if "timi" in quant_results:
        lines.append("## timi 精度结果")
        lines.append("")
        lines.append("| 配置 | 样本数 | Pass Rate | Avg Max Diff | Peak Max Diff | Mean Label Acc |")
        lines.append("| --- | ---: | ---: | ---: | ---: | ---: |")
        vendor_baseline = vendor_baselines.get("timi", {})
        vendor_summary = vendor_baseline.get("summary")
        if vendor_summary is not None:
            lines.append(_format_timi_row(vendor_baseline.get("label", "vendor_gpu_vs_pytorch_ref"), vendor_summary))
        for item in quant_results["timi"]:
            lines.append(_format_timi_row(item["quant_type"], item["compare"]["summary"]))
        lines.append("")

    if "hxx_lstm" in quant_results:
        lines.append("## hxx_lstm 精度结果")
        lines.append("")
        lines.append("| 配置 | MAE | MAPE(%) |")
        lines.append("| --- | ---: | ---: |")
        vendor_baseline = vendor_baselines.get("hxx_lstm", {})
        vendor_summary = vendor_baseline.get("summary")
        if vendor_summary is not None:
            lines.append(_format_hxx_row(vendor_baseline.get("label", "vendor_gpu_vs_model_io_base"), vendor_summary))
        for item in quant_results["hxx_lstm"]:
            lines.append(_format_hxx_row(item["quant_type"], item["compare"]["summary"]))
        lines.append("")

    lines.append("## 结论")
    lines.append("")
    if "timi" in quant_results and quant_results["timi"]:
        best_timi = min(
            quant_results["timi"],
            key=lambda item: item["compare"]["summary"]["numerical_summary"]["max_diff_mean"],
        )
        lines.append(
            f"- `timi` 以 `Avg Max Diff` 计，最佳量化配置是 `{best_timi['quant_type']}`，"
            f"平均最大差异为 {best_timi['compare']['summary']['numerical_summary']['max_diff_mean']:.6e}。"
        )
    if "hxx_lstm" in quant_results and quant_results["hxx_lstm"]:
        best_hxx = min(
            quant_results["hxx_lstm"],
            key=lambda item: item["compare"]["summary"]["MAE"],
        )
        lines.append(
            f"- `hxx_lstm` 以 `MAE` 计，最佳量化配置是 `{best_hxx['quant_type']}`，"
            f"MAE 为 {best_hxx['compare']['summary']['MAE']:.6e}。"
        )
    lines.append("")

    lines.append("## 产物目录")
    lines.append("")
    lines.append("- `work_dirs/zte/l3_precision/run_manifest.json`: 分阶段执行的 manifest，包含 hmonnx 路径和预期输出路径")
    lines.append("- `work_dirs/zte/l3_precision/<case>/<quant_type>/hmonnx/`: 量化后的 hmonnx")
    lines.append("- `work_dirs/zte/l3_precision/<case>/<quant_type>/outputs/`: hmonnx 推理输出")
    lines.append("- `work_dirs/zte/l3_precision/<case>/<quant_type>/compare/`: 对比报告")
    lines.append("- `work_dirs/zte/l3_precision/<case>/vendor_baseline/`: 对方 GPU 输出的基线报告")
    lines.append("")

    summary_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return summary_path


def write_summary_json(
    output_root: Path,
    l3_test_root: Path,
    quant_results: dict[str, list[dict[str, Any]]],
    vendor_baselines: dict[str, dict[str, Any]],
) -> Path:
    summary_path = output_root / "summary.json"
    payload = {
        "generated_at": datetime.now().isoformat(),
        "l3_test_root": str(l3_test_root),
        "vendor_baselines": vendor_baselines,
        "quant_results": quant_results,
    }
    summary_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=_json_default) + "\n", encoding="utf-8")
    return summary_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="ZTE L3 timi/hxx_lstm 量化、推理、对比一键脚本")
    parser.add_argument(
        "--cases",
        nargs="+",
        default=["timi", "hxx_lstm"],
        choices=sorted(CASE_NAMES),
        help="要执行的测试 case。",
    )
    parser.add_argument(
        "--quant-types",
        nargs="+",
        default=["w8a8_sefp", "w8a16_sefp", "w16a16_sefp"],
        help="量化类型列表。",
    )
    parser.add_argument(
        "--stages",
        nargs="+",
        default=list(PIPELINE_STAGES),
        choices=list(PIPELINE_STAGES),
        help="执行阶段，可选 export / infer / compare。默认全部执行。",
    )
    parser.add_argument(
        "--output-root",
        type=Path,
        default=REPO_ROOT / "work_dirs/zte/l3_precision",
        help="输出目录。",
    )
    parser.add_argument(
        "--l3-test-root",
        type=Path,
        default=None,
        help=f"L3 精度测试数据目录，可直接传 `L3精度测试用例` 或其上层目录；也可通过环境变量 `{L3_TEST_ROOT_ENV}` 指定。",
    )
    parser.add_argument(
        "--device",
        choices=["auto", "cpu", "cuda"],
        default="auto",
        help="HMONNX 推理设备。",
    )
    parser.add_argument(
        "--runtime-float-dtype",
        choices=["fp16", "fp32"],
        default="fp16",
        help="HMONNX 推理时 float 输入的 dtype。",
    )
    parser.add_argument(
        "--skip-vendor-baseline",
        action="store_true",
        help="跳过对方 GPU 输出基线对比。",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="覆盖已存在的 hmonnx、输出和汇总。",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_root = args.output_root.resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    l3_test_root = resolve_l3_test_root(args.l3_test_root)
    case_configs = build_case_configs(l3_test_root)

    device = _resolve_device(args.device)
    runtime_float_dtype = _get_runtime_float_dtype(args.runtime_float_dtype)
    selected_stages = set(args.stages)

    if STAGE_EXPORT in selected_stages:
        manifest_path = run_export_stage(case_configs, args.cases, args.quant_types, output_root, l3_test_root, args.overwrite)
        print(f"Run Manifest: {manifest_path}")

    if STAGE_INFER in selected_stages:
        infer_overwrite = args.overwrite and STAGE_EXPORT not in selected_stages
        for case_name in args.cases:
            case_config = case_configs[case_name]
            model_path = prepare_model_for_case(case_config, output_root, infer_overwrite)
            for quant_type in args.quant_types:
                hmonnx_path = quantize_to_hmonnx(case_config, model_path, quant_type, output_root, infer_overwrite)
                output_path, infer_stats = run_inference(
                    case_config=case_config,
                    model_path=model_path,
                    hmonnx_path=hmonnx_path,
                    output_root=output_root,
                    quant_type=quant_type,
                    device=device,
                    runtime_float_dtype=runtime_float_dtype,
                    overwrite=infer_overwrite,
                )
                print(f"Infer [{case_name}/{quant_type}]: {output_path} ({infer_stats})")

    if STAGE_COMPARE in selected_stages:
        summary_json, summary_md = run_compare_stage(
            case_configs=case_configs,
            cases=args.cases,
            quant_types=args.quant_types,
            output_root=output_root,
            l3_test_root=l3_test_root,
            device=device,
            runtime_float_dtype=runtime_float_dtype,
            overwrite=args.overwrite,
            skip_vendor_baseline=args.skip_vendor_baseline,
        )
        print(f"Summary JSON: {summary_json}")
        print(f"Summary Markdown: {summary_md}")


if __name__ == "__main__":
    main()
