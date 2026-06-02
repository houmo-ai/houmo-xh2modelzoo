"""Shared utilities for the three customer-docx eval scripts.

Backend-agnostic: prompt construction, docx I/O, summary writing.
The three eval_*.py scripts only differ in the model/backend loader.
"""
# Copyright 2025 HOUMO AI
# SPDX-License-Identifier: Apache-2.0

import json
import re
import time
from dataclasses import dataclass
from pathlib import Path
from typing import List

DEFAULT_PROMPT_MD = "/data01/home/yujy/work/xunfei/xf_data/prompt.md"
DEFAULT_DOCX_DIR = "/data01/home/yujy/work/xunfei/xf_data"
DEFAULT_OUT_ROOT = "work_dirs/customer_eval"

# === Generation config (统一后处理参数; per user spec 2026-06-02) ===
GEN_TEMPERATURE = 0.3
GEN_TOP_P = 0.9
GEN_TOP_K = 20
GEN_REPETITION_PENALTY = 1.1
GEN_MAX_NEW_TOKENS = 32768
GEN_DO_SAMPLE = True   # temperature=0.3 needs sampling; do_sample=False would be greedy


def load_system_prompt(prompt_md_path: str) -> str:
    """Extract SYSTEM_PROMPT triple-quoted literal from prompt.md.

    prompt.md is a python snippet: SYSTEM_PROMPT = \"\"\"...\"\"\". Regex the
    triple-quoted literal so we never execute untrusted file content.
    """
    text = Path(prompt_md_path).read_text(encoding="utf-8")
    m = re.search(
        r'SYSTEM_PROMPT\s*=\s*(?P<q>"""|\'\'\')(?P<body>.*?)(?P=q)',
        text,
        re.DOTALL,
    )
    if not m:
        raise ValueError(f"Could not find SYSTEM_PROMPT triple-quoted literal in {prompt_md_path}")
    return m.group("body").strip()


def extract_docx_text(docx_path: str) -> str:
    """Read paragraphs from a .docx, joining non-empty ones with newline.

    Speaker markers (<审>：/<原代>： or bare 审判长：/原代：) live inline in
    the paragraph text, so a plain join preserves them. Empty paragraphs
    are dropped to avoid blank-line runs.
    """
    from docx import Document

    doc = Document(docx_path)
    lines: List[str] = []
    for para in doc.paragraphs:
        t = para.text
        if t is not None and t.strip() != "":
            lines.append(t)
    return "\n".join(lines)


def build_messages(system_prompt: str, transcript: str) -> List[dict]:
    return [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": transcript},
    ]


# === Docx output ===

def write_docx_output(out_path: Path, *, title: str, meta: dict, body_text: str) -> None:
    """Write a docx with metadata header (paragraph) + body paragraphs.

    Format mirrors a condensed court transcript record: title + meta block
    at the top, then the generated text split into paragraphs by newline.
    """
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    # Title
    h = doc.add_heading(title, level=1)
    # Meta block
    for k, v in meta.items():
        p = doc.add_paragraph()
        run = p.add_run(f"{k}: ")
        run.bold = True
        p.add_run(str(v))
    # Body
    body_text = (body_text or "").strip("\n")
    if not body_text:
        doc.add_paragraph("[EMPTY OUTPUT]").italic = True
    else:
        for para in body_text.split("\n"):
            para = para.strip()
            if not para:
                continue
            p = doc.add_paragraph(para)
            for run in p.runs:
                run.font.size = Pt(11)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


# === Backend protocol ===

@dataclass
class GenResult:
    text: str
    in_tokens: int
    out_tokens: int
    seconds: float


class BaseBackend:
    """Subclasses implement .build_prompt_text, .count_tokens, .generate."""

    def build_prompt_text(self, messages: List[dict]) -> str:
        raise NotImplementedError

    def count_tokens(self, prompt_text: str) -> int:
        raise NotImplementedError

    def generate(self, prompt_text: str, max_new_tokens: int) -> str:
        raise NotImplementedError

    def run(self, messages: List[dict], max_new_tokens: int) -> GenResult:
        prompt_text = self.build_prompt_text(messages)
        n_in = self.count_tokens(prompt_text)
        t0 = time.time()
        out_text = self.generate(prompt_text, max_new_tokens)
        dt = time.time() - t0
        # token count on the decoded output (best-effort: re-tokenize the output)
        n_out = self.count_tokens(out_text) if out_text else 0
        return GenResult(text=out_text or "", in_tokens=n_in,
                         out_tokens=n_out, seconds=round(dt, 1))


# === Run loop ===

def discover_docx(docx_dir: str, only: str = None, max_docs: int = 0) -> List[Path]:
    files = sorted(Path(docx_dir).glob("*.docx"))
    if not files:
        raise FileNotFoundError(f"no .docx under {docx_dir}")
    if only:
        files = [f for f in files if only in f.name]
    if max_docs and max_docs > 0:
        files = files[: max_docs]
    if not files:
        raise FileNotFoundError(f"no .docx left after filtering (only={only})")
    return files


def write_summary(out_dir: Path, summary: List[dict]) -> Path:
    p = out_dir / "_summary.json"
    p.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return p
