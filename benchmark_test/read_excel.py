import openpyxl
import json

wb = openpyxl.load_workbook('Model Zoo.xlsx')
ws = wb.active

print('工作表名称:', wb.sheetnames)
print('\n前30行数据:')
for i, row in enumerate(ws.iter_rows(values_only=True), 1):
    print(row)
    if i >= 30:
        break

print('\n列名:')
for col in ws[1]:
    print(col.value)