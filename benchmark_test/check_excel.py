import openpyxl

wb = openpyxl.load_workbook('Model Zoo.xlsx')
print('工作表名称:', wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n=== 工作表: {sheet_name} ===')
    print('前10行数据:')
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        print(row)
        if i >= 10:
            break